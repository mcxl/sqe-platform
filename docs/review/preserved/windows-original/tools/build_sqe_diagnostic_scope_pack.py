from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = Path(r"LOCAL_HOME\OneDrive - AuditCo\SQE\05 Rapid Diagnostic and Scope Engagement Pack")
OUT_DIR.mkdir(parents=True, exist_ok=True)

# Simple black-and-white visual system, with no cover page.
NAVY = "000000"
BLUE = "000000"
DARK_BLUE = "000000"
LIME = "666666"
ORANGE = "555555"
PALE_BLUE = "F2F2F2"
PALE_GREEN = "E7E7E7"
PALE_ORANGE = "EDEDED"
LIGHT_GREY = "F2F2F2"
MID_GREY = "555555"
DARK_GREY = "222222"
BLACK = "000000"
WHITE = "FFFFFF"
LIGHT_BORDER = "D0D5DD"
PALE_NAVY = "E7E7E7"

A4_WIDTH = Inches(8.2677165)
A4_HEIGHT = Inches(11.6929134)
MARGIN = Inches(0.70)
PAGE_WIDTH_DXA = 9890


def set_run_font(run, *, size: float | None = None, bold: bool | None = None,
                 italic: bool | None = None, color: str = BLACK,
                 name: str = "Calibri") -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top: int = 80, bottom: int = 80,
                     start: int = 110, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("bottom", bottom),
                        ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, **edges) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, options in edges.items():
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        for key, value in options.items():
            node.set(qn(f"w:{key}"), str(value))


def set_table_geometry(table, widths_dxa: Sequence[int], indent_dxa: int = 0) -> None:
    if sum(widths_dxa) != PAGE_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {PAGE_WIDTH_DXA}: {widths_dxa}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(PAGE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    set_run_font(run, size=8.2, color=MID_GREY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    visible = OxmlElement("w:t")
    visible.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, visible, end])


def add_numbering_definition(doc: Document, fmt: str, marker: str,
                             *, left: int = 720, hanging: int = 360) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId")))
                    for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId")))
               for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), marker)
    lvl.append(lvl_text)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    lvl.append(jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    p_pr.append(ind)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(abstract_id))
    num.append(abs_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])


def configure_document(doc: Document, *, infographic: bool = False) -> None:
    sec = doc.sections[0]
    sec.page_width = A4_WIDTH
    sec.page_height = A4_HEIGHT
    sec.top_margin = Inches(0.62 if infographic else 0.64)
    sec.bottom_margin = Inches(0.60 if infographic else 0.62)
    sec.left_margin = MARGIN
    sec.right_margin = MARGIN
    sec.header_distance = Inches(0.30)
    sec.footer_distance = Inches(0.30)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.4 if not infographic else 9.2)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(5 if not infographic else 3)
    normal.paragraph_format.line_spacing = 1.08 if not infographic else 1.0

    heading_tokens = {
        1: (16, NAVY, 13, 6),
        2: (12.7, BLUE, 9, 4),
        3: (11.2, DARK_BLUE, 7, 3),
    }
    for level, (size, colour, before, after) in heading_tokens.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size if not infographic else max(size - 1, 9.5))
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(colour)
        style.paragraph_format.space_before = Pt(before if not infographic else max(before - 3, 2))
        style.paragraph_format.space_after = Pt(after if not infographic else max(after - 1, 2))
        style.paragraph_format.keep_with_next = True

    header = sec.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.paragraph_format.space_after = Pt(0)
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("Squadron Energy | Safety Risk Management Internal Audit")
    set_run_font(hr, size=8.1 if not infographic else 7.5, bold=True, color=BLACK)

    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.paragraph_format.space_after = Pt(0)
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("Confidential - Working Draft | Page ")
    set_run_font(fr, size=8.1 if not infographic else 7.5, color=MID_GREY)
    add_page_field(fp)

    props = doc.core_properties
    props.author = "AuditCo"
    props.last_modified_by = "AuditCo"
    props.subject = "Squadron Energy safety risk management internal audit"


def add_para(doc: Document, text: str = "", *, size: float = 10.4,
             bold: bool = False, italic: bool = False, color: str = BLACK,
             before: float = 0, after: float = 5,
             align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing: float = 1.08,
             keep: bool = False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.keep_with_next = keep
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_rich_para(doc: Document, segments: Sequence[tuple[str, dict]], *,
                  size: float = 10.4, before: float = 0, after: float = 5,
                  line_spacing: float = 1.08, keep: bool = False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.keep_with_next = keep
    for text, kwargs in segments:
        r = p.add_run(text)
        set_run_font(r, size=kwargs.get("size", size),
                     bold=kwargs.get("bold"), italic=kwargs.get("italic"),
                     color=kwargs.get("color", BLACK))
    return p


def add_heading(doc: Document, text: str, level: int = 1, *, page_break: bool = False):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.page_break_before = page_break
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_bullets(doc: Document, items: Iterable[str], *, bullet_id: int | None = None,
                size: float = 10.4, after: float = 3.5):
    if bullet_id is None:
        bullet_id = add_numbering_definition(doc, "bullet", "•", left=620, hanging=320)
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.08
        apply_num(p, bullet_id)
        r = p.add_run(item)
        set_run_font(r, size=size, color=BLACK)
    return bullet_id


def add_numbered(doc: Document, items: Iterable[str], *, size: float = 10.2,
                 after: float = 3.0):
    num_id = add_numbering_definition(doc, "decimal", "%1.", left=620, hanging=320)
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.06
        apply_num(p, num_id)
        r = p.add_run(item)
        set_run_font(r, size=size, color=BLACK)
    return num_id


def write_cell(cell, value, *, size: float = 8.5, bold: bool = False,
               color: str = BLACK, align=WD_ALIGN_PARAGRAPH.LEFT,
               line_spacing: float = 1.02) -> None:
    cell.text = ""
    values = value if isinstance(value, list) else [str(value)]
    for index, text in enumerate(values):
        p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1 if len(values) > 1 else 0)
        p.paragraph_format.line_spacing = line_spacing
        r = p.add_run(str(text))
        set_run_font(r, size=size, bold=bold, color=color)


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence],
              widths: Sequence[int], *, font_size: float = 8.4,
              header_fill: str = NAVY, first_col_fill: str | None = PALE_NAVY,
              alignments: Sequence | None = None, compact: bool = False):
    table = doc.add_table(rows=1, cols=len(headers))
    for header, cell in zip(headers, table.rows[0].cells):
        shade_cell(cell, header_fill)
        write_cell(cell, header, size=font_size, bold=True, color=WHITE,
                   align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0)
        set_cell_borders(
            cell,
            top={"val": "single", "sz": "4", "color": header_fill},
            bottom={"val": "single", "sz": "4", "color": header_fill},
            start={"val": "single", "sz": "4", "color": WHITE},
            end={"val": "single", "sz": "4", "color": WHITE},
        )
    repeat_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for col_index, value in enumerate(values):
            alignment = alignments[col_index] if alignments else WD_ALIGN_PARAGRAPH.LEFT
            write_cell(row.cells[col_index], value, size=font_size,
                       align=alignment, line_spacing=1.0 if compact else 1.03)
            fill = LIGHT_GREY if row_index % 2 == 0 else WHITE
            if first_col_fill and col_index == 0:
                fill = first_col_fill
            shade_cell(row.cells[col_index], fill)
            set_cell_borders(
                row.cells[col_index],
                top={"val": "single", "sz": "3", "color": LIGHT_BORDER},
                bottom={"val": "single", "sz": "3", "color": LIGHT_BORDER},
                start={"val": "single", "sz": "3", "color": LIGHT_BORDER},
                end={"val": "single", "sz": "3", "color": LIGHT_BORDER},
            )
    set_table_geometry(table, widths)
    add_para(doc, after=3)
    return table


def add_callout(doc: Document, label: str, text: str, *, fill: str = PALE_BLUE,
                accent: str = BLUE, size: float = 10.1):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_borders(
        cell,
        top={"val": "nil"}, bottom={"val": "nil"}, end={"val": "nil"},
        start={"val": "single", "sz": "18", "color": accent},
    )
    set_cell_margins(cell, top=110, bottom=110, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(label)
    set_run_font(r, size=8.8, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.05
    r2 = p2.add_run(text)
    set_run_font(r2, size=size, color=BLACK)
    set_table_geometry(table, [PAGE_WIDTH_DXA])
    add_para(doc, after=3)
    return table


def add_title_block(doc: Document, *, title: str, subtitle: str,
                    purpose: str, meta: Sequence[tuple[str, str]],
                    kicker: str = "AuditCo") -> None:
    add_para(doc, kicker, size=10.2, bold=True, color=BLACK, before=1, after=2, keep=True)
    add_para(doc, title, size=22.5, bold=True, color=NAVY, after=3,
             line_spacing=0.98, keep=True)
    add_para(doc, subtitle, size=12.2, bold=True, color=DARK_BLUE, after=8,
             line_spacing=1.0, keep=True)
    add_callout(doc, "Purpose", purpose, fill=PALE_BLUE, accent=BLUE, size=10.1)
    table = doc.add_table(rows=len(meta), cols=2)
    for index, (label, value) in enumerate(meta):
        write_cell(table.rows[index].cells[0], label, size=8.8, bold=True,
                   color=NAVY)
        write_cell(table.rows[index].cells[1], value, size=8.8, color=BLACK)
        shade_cell(table.rows[index].cells[0], LIGHT_GREY)
        shade_cell(table.rows[index].cells[1], WHITE)
        prevent_row_split(table.rows[index])
        for cell in table.rows[index].cells:
            set_cell_borders(cell, top={"val": "nil"}, bottom={"val": "nil"},
                             start={"val": "nil"}, end={"val": "nil"})
    set_table_geometry(table, [2150, PAGE_WIDTH_DXA - 2150])
    add_para(doc, after=4)


def add_scope_focus_table(doc: Document, scopes: Sequence[dict]) -> None:
    rows = []
    for item in scopes:
        rows.append((item["name"], item["primary_question"], item["focus_line"]))
    add_table(doc, ["Scope Area", "Primary Question", "Focus"], rows,
              [2500, 5000, 2390], font_size=8.0, first_col_fill=PALE_GREEN,
              compact=True)


def add_scope_timing_summary_table(doc: Document, scopes: Sequence[dict]) -> None:
    rows = []
    for item in scopes:
        rows.append((
            item["name"],
            item["focus_line"] + ".",
            [
                "3 working days: " + item["timing"]["focused"].split(" - ", 1)[1],
                "4 working days: " + item["timing"]["standard"].split(" - ", 1)[1],
                "5 working days: " + item["timing"]["complex"].split(" - ", 1)[1],
            ],
        ))
    add_table(doc, ["Scope Area", "Summary Breakdown", "Indicative Timing"], rows,
              [2500, 3300, 4090], font_size=7.7, first_col_fill=PALE_GREEN,
              compact=True)


def add_standard_scope_sequence(doc: Document, *, compact: bool = False) -> None:
    rows = [
        ("Activation", "Confirm written election, independence, fee basis, scope and start conditions.", "Gate 0 - Start readiness"),
        ("Scope Planning", "Confirm criteria, assets, lifecycle stages, evidence, people, walkthroughs, sample and report needs.", "Gate 1 - Plan approved"),
        ("Evidence Review", "Build the criteria-to-control matrix and record gaps, contradictions and limitations.", "Connected evidence review"),
        ("Detailed Fieldwork", "Complete interviews, walkthroughs, document review, sample review and bounded checks where relevant.", "Gate 2 - Fieldwork complete"),
        ("Assessment And Challenge", "Apply MATE, complete CONTRA challenge and develop evidence-supported observations and actions.", "Gate 3 - Proceed to draft"),
        ("Draft And Factual Review", "Prepare draft report and action tracker; complete independent partner review; receive one consolidated client response.", "Gate 4 - Facts checked"),
        ("Finalisation And Close", "Resolve supported corrections, confirm actions and issue the final report; transfer the tracker and record the next election.", "Gate 5 - Final report issued"),
    ]
    add_table(doc, ["Stage", "AuditCo Work", "Gate Or Output"], rows,
              [1800, 5700, 2390], font_size=8.0, first_col_fill=PALE_NAVY,
              compact=compact)


SCOPES = [
    {
        "number": 1,
        "name": "Governance And Oversight",
        "file": "03_Governance_And_Oversight_Separate_Engagement_Plan",
        "primary_question": "Does governance provide clear authority, information, oversight and escalation for critical safety risk?",
        "focus_line": "Authority, reporting, oversight and escalation",
        "objective": "Assess whether governance provides clear authority, reliable safety information, accountable decision-making, oversight of critical risk and timely escalation.",
        "questions": [
            "Are accountabilities, delegations and decision authorities clear?",
            "Do executives and committees receive reliable safety information with defined measures and thresholds?",
            "Are critical controls, significant events and overdue actions visible to the right decision-maker?",
            "Are reporting triggers and escalation paths defined and connected to projects, operations and contractors?",
            "Are obligations monitoring, training oversight and assurance activity connected to governance?",
        ],
        "evidence": [
            "Organisation charts, governance maps and delegations",
            "Committee and management terms of reference, papers, minutes and action trackers",
            "Safety dashboards, indicator definitions and reporting protocols",
            "Obligations monitoring, training oversight and assurance plans",
            "Project Kraken governance and decision records",
        ],
        "procedures": [
            "Map governance bodies, roles, decisions, information flows and escalation routes.",
            "Trace selected indicators from source data to management and committee decisions.",
            "Review whether critical-control information has definitions, owners, frequency and thresholds.",
            "Trace a significant event, overdue action or accepted risk to the accountable decision-maker.",
            "Test the connection between governance reporting and project, operational and contractor information.",
            "Apply MATE and complete CONTRA challenge of the proposed conclusion.",
        ],
        "interviews": [
            "Internal Audit sponsor or accountable executive",
            "Senior safety or risk executive",
            "Corporate Risk and Compliance Committee secretariat",
            "Operations Manager, project or construction leader",
            "Risk and assurance owner and selected control owners",
        ],
        "walkthroughs": [
            "Committee report from source data to decision",
            "Significant event or overdue action escalation",
            "Governance response to a material risk change",
            "Selected indicators, decisions and action closures",
        ],
        "outputs": [
            "Governance accountability map",
            "Information and escalation trace",
            "Evidence-supported findings on governance design",
            "Evidence-confidence and limitation statement",
            "Action tracker with accountable owners and dependencies",
        ],
        "timing_drivers": [
            "Number and maturity of governance bodies and reporting streams",
            "Volume of committee packs, indicators and action records",
            "Availability of decision owners and the required reporting audience",
            "Need for project, operational or contractor interface checks",
        ],
        "timing": {
            "focused": "3 working days - one governance structure and a limited reporting sample.",
            "standard": "4 working days - standard governance and committee review.",
            "complex": "5 working days - multiple committees, projects, contractors or complex reporting flows.",
        },
    },
    {
        "number": 2,
        "name": "Safety Management Framework",
        "file": "04_Safety_Management_Framework_Separate_Engagement_Plan",
        "primary_question": "Does the framework connect requirements, roles, controls, competence and assurance across the selected lifecycle?",
        "focus_line": "Framework coherence, competence, change and assurance",
        "objective": "Assess whether the safety management framework connects requirements, roles, controls, competence, document control and assurance across the selected lifecycle.",
        "questions": [
            "Are policies, standards and procedures coherent and connected to risk?",
            "Are current, transition and intended target states understood and governed?",
            "Does Project Kraken have clear ownership, dependencies and decision triggers?",
            "Are the Golden Rules controlled, communicated and supported by competence requirements?",
            "Do framework changes trigger risk, training, assurance and reporting updates?",
        ],
        "evidence": [
            "Framework hierarchy, document map, policies, standards and procedures",
            "Project Kraken current, transition and target-state records",
            "Golden Rules, rollout, communications and exception records",
            "Training, competence and authorisation matrices",
            "Document approval, change and assurance records",
            "Project, asset and operational interface records",
        ],
        "procedures": [
            "Map the framework hierarchy and identify control ownership.",
            "Review consistency across corporate, project and operational documents.",
            "Trace one material risk from framework requirement to expected control and evidence.",
            "Review Project Kraken governance, dependencies, transition conditions and target-state decisions.",
            "Review Golden Rules ownership, approval, communication, competence and exception handling.",
            "Assess whether framework changes trigger risk, training, assurance and reporting updates; apply MATE and CONTRA.",
        ],
        "interviews": [
            "Framework owner and Project Kraken owner",
            "Golden Rules owner",
            "Training and competency owner",
            "Operations Manager, project or construction leader",
            "Assurance or internal audit representative",
        ],
        "walkthroughs": [
            "Framework change from proposal to approval and release",
            "Golden Rules rollout and competence evidence",
            "Project Kraken decision or transition gate",
            "Selected assurance and document-change records",
        ],
        "outputs": [
            "Framework hierarchy and control map",
            "Current-to-target governance observations",
            "Findings on coherence, rollout, competence and assurance",
            "Evidence-confidence and limitation statement",
            "Action tracker with owners and dependencies",
        ],
        "timing_drivers": [
            "Size and complexity of the framework hierarchy",
            "Number of lifecycle states, projects and operational interfaces",
            "Volume of Golden Rules, training and competence evidence",
            "Extent of Project Kraken transition and target-state evidence",
        ],
        "timing": {
            "focused": "3 working days - one defined framework component.",
            "standard": "4 working days - standard framework review.",
            "complex": "5 working days - framework review including Project Kraken, Golden Rules, competence and transition evidence.",
        },
    },
    {
        "number": 3,
        "name": "Risk Register Integration",
        "file": "05_Risk_Register_Integration_Separate_Engagement_Plan",
        "primary_question": "Do risk registers connect hazards, controls, owners, triggers, actions and assurance across corporate, project and operational levels?",
        "focus_line": "Risk ownership, controls, triggers and assurance",
        "objective": "Assess whether material safety risks and critical controls are consistently identified, owned, reviewed, escalated and connected across corporate, project, asset and operational risk records.",
        "questions": [
            "Do risk registers capture high-consequence hazards and material changes?",
            "Are controls, accountable owners and acceptance authorities clear?",
            "Do risk ratings and acceptance decisions use consistent criteria?",
            "Are review triggers linked to change, events and lifecycle stages?",
            "Do risks connect to critical controls, assurance, actions and executive reporting?",
        ],
        "evidence": [
            "Risk management policy, criteria and risk appetite or acceptance rules",
            "Corporate, project, asset and operational risk registers",
            "Critical-control registers and control-verification records",
            "Risk acceptance, escalation and committee reporting records",
            "Change, management-of-change, incident and event records",
            "Assurance plans and completed reviews",
        ],
        "procedures": [
            "Map the relationship between corporate, project, asset and operational risk records.",
            "Trace selected high-consequence risks to controls, owners, evidence and assurance.",
            "Review risk review triggers for change, incidents, new assets and lifecycle movement.",
            "Test whether risk acceptance and escalation reach the correct authority.",
            "Trace a selected event or change through the risk-register update process.",
            "Compare critical-control records with risk-register control descriptions; apply MATE and CONTRA.",
        ],
        "interviews": [
            "Risk owner and safety risk manager",
            "Project risk manager",
            "Operations Manager",
            "Critical-control owner",
            "Incident or assurance owner and selected executive representative",
        ],
        "walkthroughs": [
            "Risk-register change after a material trigger",
            "High-consequence risk to critical-control verification",
            "Event or change to risk-register update",
            "Selected risk acceptance and escalation decisions",
        ],
        "outputs": [
            "Risk relationship map",
            "Criteria-to-control and risk-to-assurance trace",
            "Findings on integration, ownership, triggers and escalation",
            "Evidence-confidence and limitation statement",
            "Action tracker with control and risk owners",
        ],
        "timing_drivers": [
            "Number of registers and organisational levels in scope",
            "Number of selected high-consequence risks and controls",
            "Volume of change, incident and risk-acceptance records",
            "Need to reconcile differing risk taxonomies or criteria",
        ],
        "timing": {
            "focused": "3 working days - one risk system or a limited risk sample.",
            "standard": "4 working days - standard multi-level risk review.",
            "complex": "5 working days - multiple registers, assets, projects or risk taxonomies.",
        },
    },
    {
        "number": 4,
        "name": "Incident Learning",
        "file": "06_Incident_Learning_Separate_Engagement_Plan",
        "primary_question": "Do incidents and high-potential events lead to sound investigation, learning, action and risk improvement?",
        "focus_line": "Investigation quality, learning and action verification",
        "objective": "Assess whether significant incidents and high-potential events lead to sound investigation, system learning, owned actions, risk improvement and governance visibility.",
        "questions": [
            "Are events identified, classified and notified consistently?",
            "Do escalation thresholds match event significance and retained duties?",
            "Do investigations consider system, control, work-design and human factors?",
            "Are actions specific, risk-based, owned and verified?",
            "Do lessons update risks, controls, training and contractor oversight?",
            "Does governance identify recurring themes and unresolved exposure?",
        ],
        "evidence": [
            "Incident and high-potential event procedures",
            "Event registers, notification records and escalation decisions",
            "Investigation reports, terms of reference and causal-analysis records",
            "Action trackers, closure evidence and verification records",
            "Risk-register updates, lessons-learned communications and training",
            "Committee and management event reporting",
        ],
        "procedures": [
            "Trace selected events from notification to close-out.",
            "Review classification, notification and escalation decisions.",
            "Assess investigation scope, evidence, contrary evidence and causal analysis.",
            "Review whether actions address system, control and work-design contributors.",
            "Trace selected lessons into risk registers, standards, training and contractor controls.",
            "Review recurring themes and unresolved action exposure; apply MATE and CONTRA.",
        ],
        "interviews": [
            "Incident-learning owner and investigation lead",
            "Senior safety or risk executive",
            "Operations Manager and project or construction leader",
            "Selected action owners",
            "Contractor-governance owner where a contractor event is selected",
        ],
        "walkthroughs": [
            "High-potential event from notification to learning",
            "Action from approval to closure verification",
            "Recurring theme to risk or control change",
            "Selected overdue or reopened action",
        ],
        "outputs": [
            "Incident-learning pathway map",
            "Event-to-action and action-to-risk trace",
            "Findings on investigation, learning and escalation design",
            "Evidence-confidence and limitation statement",
            "Action tracker with verification requirements",
        ],
        "timing_drivers": [
            "Number and complexity of event files selected",
            "Depth of investigation and causal-analysis records",
            "Number of action owners and closure-verification points",
            "Need for specialist, project or contractor event context",
        ],
        "timing": {
            "focused": "3 working days - one focused event pathway or a small sample.",
            "standard": "4 working days - standard event and action review.",
            "complex": "5 working days - multiple events, complex investigations or recurring-theme analysis.",
        },
    },
    {
        "number": 5,
        "name": "Third-Party Governance",
        "file": "07_Third-Party_Governance_Separate_Engagement_Plan",
        "primary_question": "Does the contractor lifecycle protect Squadron Energy's retained duties and provide suitable oversight of critical safety risks?",
        "focus_line": "Contractor lifecycle, critical controls and retained duties",
        "objective": "Assess whether the contractor lifecycle provides risk-proportionate selection, contract requirements, mobilisation, oversight, critical-control verification, deviation control, handover and close-out.",
        "questions": [
            "Are safety requirements set during prequalification and procurement?",
            "Do executed contracts contain approved safety clauses and clear interfaces?",
            "Are mobilisation, consultation, cooperation and information-sharing requirements clear?",
            "Does oversight use suitable leading and lagging information?",
            "Are critical controls, deviations and high-potential events verified and escalated?",
            "Are handover and close-out responsibilities clear?",
        ],
        "evidence": [
            "Contractor governance standard, prequalification and procurement criteria",
            "Approved contract clause baseline",
            "Selected executed contracts and scopes of work",
            "Mobilisation, consultation and oversight plans",
            "Contractor dashboards, critical-control verification and performance reports",
            "Deviation approvals, contractor events, action records, handover and close-out records",
        ],
        "procedures": [
            "Map the contractor lifecycle from selection to close-out.",
            "Compare selected executed contracts with the approved safety-clause baseline.",
            "Review retained duties, consultation, cooperation, information-sharing and oversight interfaces.",
            "Review critical-control verification, leading and lagging information and escalation.",
            "Review deviations, approvals, safeguards and expiry or close-out controls.",
            "Trace a selected contractor event and review handover; apply MATE and CONTRA.",
        ],
        "interviews": [
            "Procurement or prequalification owner",
            "Contract owner and contractor-governance owner",
            "Project or construction leader",
            "Operations Manager",
            "Safety or risk executive and authorised principal contractor representative where appropriate",
        ],
        "walkthroughs": [
            "Prequalification or procurement decision",
            "Contract clause approval and executed contract comparison",
            "Mobilisation and oversight plan",
            "Critical-control verification and deviation escalation",
            "Handover or close-out record",
        ],
        "outputs": [
            "Contractor lifecycle and retained-duty map",
            "Contract baseline and deviation analysis",
            "Findings on oversight, critical controls, escalation and close-out",
            "Separate bounded contractor evidence statement",
            "Evidence-confidence and limitation statement and action tracker",
        ],
        "timing_drivers": [
            "Contractor population, selected contracts and jurisdictions",
            "Site or project coverage and the required sample depth",
            "Availability of procurement, contract and oversight records",
            "Need for technical or principal-contractor input",
        ],
        "timing": {
            "focused": "3 working days - one contractor and a narrow lifecycle review.",
            "standard": "4 working days - standard contractor lifecycle and contract sample.",
            "complex": "5 working days - multiple contractors, projects, jurisdictions or technical interfaces.",
        },
    },
    {
        "number": 6,
        "name": "Safety In Design And Handover",
        "file": "08_Safety_In_Design_And_Handover_Separate_Engagement_Plan",
        "primary_question": "Does the design and handover process identify, control and transfer residual safety risk into operations?",
        "focus_line": "Design risk, residual risk and operational transfer",
        "objective": "Assess whether design and handover processes identify foreseeable hazards, control residual risk, govern changes and transfer usable safety information into operations.",
        "questions": [
            "Are foreseeable hazards identified early and addressed through design?",
            "Are elimination, substitution and risk-reduction decisions recorded?",
            "Are residual risks owned, accepted, controlled and transferred?",
            "Do design reviews and changes have clear triggers and approval routes?",
            "Is commissioning information complete and controlled?",
            "Can the receiving operational team find, understand and use the safety information?",
            "Are duty-holder interfaces and retained responsibilities clear?",
        ],
        "evidence": [
            "Safety-in-design process and criteria",
            "Design risk registers, hazard logs, design reviews and action trackers",
            "Drawings, specifications and approved design changes",
            "Manufacturer and supplier information",
            "Commissioning, inspection, testing and operational-readiness records",
            "Residual-risk registers, as-built and isolation information, emergency requirements and health and safety file",
        ],
        "procedures": [
            "Trace selected hazards from identification to control and residual-risk acceptance.",
            "Review design review points, accountabilities, inputs, outputs and escalation triggers.",
            "Assess whether change control reopens design risk decisions when conditions change.",
            "Review residual-risk ownership, safeguards and transfer requirements.",
            "Trace commissioning and operational-readiness evidence into handover records.",
            "Assess whether the receiving operational team can use the information; apply MATE and CONTRA.",
        ],
        "interviews": [
            "Design manager and safety-in-design owner",
            "Project or construction leader",
            "Commissioning manager",
            "Operations Manager",
            "Document-control or handover owner and technical specialist where appropriate",
        ],
        "walkthroughs": [
            "Design change from trigger to approval",
            "Residual risk from design record to operational control",
            "Commissioning and operational-readiness gate",
            "Handover information from project to operations",
            "Selected as-built, isolation and emergency records",
        ],
        "outputs": [
            "Design-to-handover risk trace",
            "Residual-risk and information-transfer observations",
            "Findings on design review, change and operational readiness",
            "Evidence-confidence and limitation statement",
            "Action tracker with design, project and operations owners",
        ],
        "timing_drivers": [
            "Selected projects or assets and lifecycle stage",
            "Volume and maturity of design, change and commissioning records",
            "Handover and operational-readiness evidence availability",
            "Need for technical specialist or contractor input",
        ],
        "timing": {
            "focused": "3 working days - one design-to-handover risk thread.",
            "standard": "4 working days - standard project or asset review.",
            "complex": "5 working days - multiple projects, assets, design changes or technical interfaces.",
        },
    },
]


def build_eoi() -> Path:
    doc = Document()
    configure_document(doc)
    add_title_block(
        doc,
        title="Rapid Five-Day Diagnostic And Elected Scope Engagements",
        subtitle="Safety Risk Management Internal Audit - Expression Of Interest Response",
        purpose="Provide Squadron Energy with a rapid, decision-useful introduction to the six scope areas, identify the priority area for detailed focus and support separate billed engagements activated in the order and timing chosen by Squadron Energy.",
        meta=[
            ("Prepared for", "Squadron Energy"),
            ("Prepared by", "AuditCo Australia"),
            ("Date", "11 August 2026"),
            ("Status", "Revised working draft for scope, commercial and independence approval"),
        ],
    )
    add_heading(doc, "1. Recommended Delivery Position")
    add_para(doc, "AuditCo proposes a standalone rapid diagnostic lasting five working days and introducing all six scope areas at a high level. The diagnostic is designed to establish the cross-scope evidence position, compare priority risk questions and identify the scope area that needs detailed focus.")
    add_para(doc, "After the diagnostic, Squadron Energy may elect a separate billed engagement for the recommended scope area, another scope area, several scope areas or no further work. Each elected scope area will have its own written scope confirmation, fee, timing, evidence request, work plan, report and action tracker.")
    add_callout(doc, "Commercial control", "No later scope engagement starts automatically. AuditCo confirms the scope, procedures, evidence, fee and expected issue point after Squadron Energy makes an election.", fill=PALE_GREEN, accent=LIME)

    add_heading(doc, "2. Response To The Submission Requirements")
    add_table(doc, ["Requirement", "AuditCo Response", "Completion Position"], [
        ("Capability And Experience", "AuditCo brings safety governance, management-system, contractor-assurance and construction experience. Direct renewable electrical and commissioning expertise will be provided through the nominated technical specialist.", "Team names, role allocation and approved experience references are to be confirmed."),
        ("Proposed Team", "Engagement Partner, Lead Internal Auditor, Safety Governance Specialist, Audit Manager or Analyst and technical specialists selected for the authorised assets and lifecycle stages.", "Names, availability and approximate effort are to be confirmed."),
        ("Approach And Timing", "The work starts with a five-working-day diagnostic across the six scope areas at an introductory level. It identifies the priority scope for detailed focus. Separate scope engagements are available for the six scope areas, and each starts only when Squadron Energy elects to activate it. Each separate engagement is expected to take 3-5 working days, with the exact duration confirmed before activation.", "Subject to access, evidence readiness and client availability."),
        ("Fees", "The diagnostic is a fixed fee of $6,000 excluding goods and services tax. Each later scope engagement has its own fixed or capped fee agreed before activation.", "Payment milestones and any agreed credit for directly reusable diagnostic work are to be confirmed."),
        ("Independence", "AuditCo will complete and approve an independence assessment before issue, recording prior services, actual or perceived conflicts and safeguards.", "Final clearance is to be confirmed."),
    ], [1900, 5200, 2790], font_size=7.9, compact=True)

    add_heading(doc, "3. The Six Scope Areas")
    add_scope_focus_table(doc, SCOPES)
    add_heading(doc, "Indicative Timing By Scope Area", 2)
    add_scope_timing_summary_table(doc, SCOPES)
    add_callout(doc, "Client-facing timing position", "Each separately elected scope engagement is expected to take 3-5 working days. Three days is a focused review, four days is the standard basis and five days is used for greater complexity. AuditCo confirms the final timing and fee before activation.", fill=PALE_ORANGE, accent=ORANGE, size=9.3)
    add_para(doc, "The diagnostic considers each scope area at a high level. The detailed scope, assets, entities, jurisdictions and lifecycle stages will be confirmed in the separate scope confirmation. The list above is a menu of separately activatable work packages, not a commitment to complete all areas.", size=9.5, italic=True, color=MID_GREY)

    add_heading(doc, "4. Assurance Method And Boundary")
    add_para(doc, "AuditCo will organise the design assessment around four connected questions:")
    add_table(doc, ["Dimension", "Design Question"], [
        ("Mandate", "Is the requirement documented, approved and connected to relevant obligations and risk?"),
        ("Accountability", "Is there a named owner with appropriate authority, delegation and oversight?"),
        ("Trigger", "Are the events, thresholds and change conditions that activate or review the control clear?"),
        ("Escalation", "Are reporting thresholds, decision paths and committee information requirements timely and unambiguous?"),
    ], [1900, 7990], font_size=8.4, first_col_fill=PALE_GREEN)
    add_para(doc, "The review will assess control design and limited evidence of implementation only to the extent needed to corroborate the design assessment. It is not a legal-compliance audit and does not provide sustained operating-effectiveness testing, engineering certification, site compliance certification or assurance over contractor performance across the portfolio.")
    add_callout(doc, "Renewable-energy risk lens", "Where relevant to the selected assets and lifecycle stages, the criteria and sample may consider high-voltage energisation, arc flash, isolation, stored energy, simultaneous operations, battery hazards, working at height, fatigue and lone work. These are risk-based criteria and do not predetermine a finding.", fill=PALE_ORANGE, accent=ORANGE, size=9.7)

    add_heading(doc, "5. Rapid Five-Day Diagnostic")
    add_table(doc, ["Day", "Focus", "AuditCo Work", "Squadron Energy Input", "Daily Output"], [
        ("Day 1", "Mobilise And Map", "Confirm the six scope areas, diagnostic boundary, key assets and lifecycle stages, criteria, priority risk questions, evidence index, people and communication route. Issue cross-scope evidence request.", "Confirm the diagnostic boundary, sponsor, coordinator, scope owners, available evidence and interview availability.", "Approved diagnostic plan and cross-scope evidence request."),
        ("Day 2", "Review Evidence And Criteria", "Review available framework, obligations, risks, controls, roles, reporting and selected records for each scope area. Build a cross-scope readiness matrix.", "Provide controlled documents or an evidence index and explain missing or superseded records.", "Cross-scope evidence and readiness matrix."),
        ("Day 3", "Targeted Interviews And Validation", "Complete priority interviews across scope owners and, where useful, one short walkthrough of the most material risk pathway. Validate the highest-priority questions.", "Provide attendees, access and process information; identify contrary evidence and constraints.", "Validated priority questions and updated evidence log."),
        ("Day 4", "Prioritise And Challenge", "Compare scope areas by risk exposure, evidence position, apparent design concern, change and management priority. Apply high-level MATE and CONTRA challenge.", "Review clarification questions and provide missing records or explain limitations.", "Priority scope recommendation and decision options."),
        ("Day 5", "Brief And Election Pack", "Issue diagnostic brief with recommended scope, rationale, evidence limitations, suggested procedures and separate engagement proposal.", "Confirm factual corrections and decide whether to request a separate scope proposal or wait.", "Diagnostic brief and client election pack."),
    ], [850, 1850, 3300, 2500, 1390], font_size=7.45, compact=True)
    add_callout(doc, "Diagnostic boundary", "The diagnostic does not complete any scope area, provide a final Control Design Adequacy rating, test sustained operating effectiveness, certify legal or technical compliance, or start a separate scope engagement without written authorisation.", fill=PALE_BLUE, accent=BLUE, size=9.5)

    add_heading(doc, "6. Separate Scope Engagement Activation")
    add_table(doc, ["Client Decision", "AuditCo Action", "Output"], [
        ("Elect the recommended scope area", "Reuse directly relevant evidence where confirmed and issue a scope confirmation for the detailed work.", "Separate fee, timing, evidence request and scope work plan."),
        ("Elect another scope area", "Review the chosen scope assumptions and issue the relevant scope confirmation.", "New scope-specific engagement basis."),
        ("Elect more than one scope area", "Prepare separate confirmations so each scope remains separately controlled and billed.", "Separate scope records and independent timing for each engagement."),
        ("Pause or close", "Record the decision and update the programme scope tracker.", "No further work until a new election is made."),
    ], [2400, 4800, 2690], font_size=8.0, compact=True)
    add_para(doc, "Each separate scope engagement is expected to be delivered in 3-5 working days. The selected duration depends on scope complexity, evidence readiness, interview and walkthrough requirements, sample depth, technical specialist input and reporting needs. AuditCo will confirm timing and the fixed or capped fee before the selected scope starts.")

    add_heading(doc, "7. Outputs And Governance")
    add_table(doc, ["Deliverable", "Purpose"], [
        ("Evidence-Readiness View", "Identifies available, missing, inconsistent or inaccessible evidence and its effect on the planned work."),
        ("Diagnostic Brief", "Summarises cross-scope evidence readiness, priority questions, apparent gaps and the recommended next scope area."),
        ("Scope-Specific Final Report", "Provides the final Control Design Adequacy conclusion for one authorised scope area under a separate engagement."),
        ("Action Tracker", "Records agreed actions, owners, due dates, priorities and dependencies."),
        ("Programme Scope Tracker", "Shows elected, active, paused, completed and remaining scope areas without creating an obligation to activate any scope."),
        ("Committee Presentation", "Provides a concise decision-focused briefing where separately agreed."),
    ], [2900, 6990], font_size=8.2, first_col_fill=PALE_NAVY)
    add_callout(doc, "Operating-effectiveness statement", "Operational Effectiveness will be shown as Not Assessed unless a separate scope expressly authorises broader testing. The reports will not provide legal advice, engineering certification or contractor-performance assurance.", fill=PALE_ORANGE, accent=ORANGE, size=9.5)

    add_heading(doc, "8. Team And Relevant Experience")
    add_table(doc, ["Role", "Required Capability", "Responsibility"], [
        ("Engagement Partner", "Senior internal audit and safety-governance leadership.", "Independence, quality, significant judgements and final conclusion."),
        ("Lead Internal Auditor - Alan Richardson", "Work health and safety, project management, construction, training and management-system audit capability.", "Planning, criteria mapping, interviews, evidence assessment, findings and reporting."),
        ("Safety Governance Specialist", "Executive governance, obligations, risk and assurance experience.", "Challenge governance information, accountability, indicators and escalation design."),
        ("Audit Manager Or Analyst", "Internal audit workpaper, analysis and reporting capability.", "Evidence control, sampling records, analysis, quality checks and action tracking."),
        ("Technical Specialist", "Direct high-voltage electrical safety and renewable-project commissioning experience where relevant.", "Challenge asset-specific criteria for energisation, isolation, stored energy, simultaneous operations and handover."),
    ], [2100, 3900, 3890], font_size=7.9, compact=True)
    add_para(doc, "AuditCo's experience includes management-system certification and auditing, internal audit work, assurance associated with a Western Australian iron ore rail network, energy-efficiency compliance and construction quality engagements. These capabilities support governance, contractor and evidence assessment, but do not replace direct renewable electrical and commissioning capability where it is required.")

    add_heading(doc, "9. Fees, Independence And Acceptance")
    add_table(doc, ["Service", "Fee Basis", "Fee"], [
        ("Rapid Five-Day Diagnostic", "Fixed fee", "$6,000 excluding goods and services tax"),
        ("Separate Scope Engagement - One Area", "Fixed or capped fee agreed before activation", "To be confirmed for the elected scope"),
        ("Additional Services", "Agreed hourly rate or separately approved fixed fee", "$135 per hour excluding goods and services tax"),
    ], [3000, 3800, 3090], font_size=8.3, compact=True)
    add_bullets(doc, [
        "Fees assume predominantly remote delivery. Travel or other disbursements are incurred only with prior approval and charged at cost.",
        "Each separate scope engagement requires written approval before commencement. Variations to scope, evidence, sample, timing, deliverables or fee require written agreement before affected work proceeds.",
        "AuditCo will not state that it is independent until the approval-controlled independence assessment has been completed.",
        "Acceptance of the diagnostic does not commit Squadron Energy to activate any later scope engagement.",
    ], size=9.5, after=2.5)
    add_heading(doc, "10. Contacts")
    add_table(doc, ["Contact", "Details"], [
        ("Squadron Energy", "Michael Worrall | +61 448 253 035 | michael.worrall@squadronenergy.com"),
        ("AuditCo", "Chris Dickson, Executive General Manager - Global | +61 477 889 319 | chris.dickson@auditco.com"),
    ], [2100, 7790], font_size=8.4, first_col_fill=PALE_NAVY, compact=True)
    add_para(doc, "Prepared by Alan Richardson. This revised EOI remains a working draft until the scope assumptions, team, independence clearance, commercial approvals and authorised signatories are confirmed.", size=9.3, italic=True, color=MID_GREY, before=3, after=0)

    path = OUT_DIR / "01_AuditCo_Squadron_Energy_EOI_Rapid_Diagnostic_and_Elected_Scope_Engagements.docx"
    doc.save(path)
    return path


def build_master_plan() -> Path:
    doc = Document()
    configure_document(doc)
    add_title_block(
        doc,
        title="Rapid Diagnostic And Scope Engagement Delivery Plan",
        subtitle="Five-day diagnostic plus separately activated scope work packages",
        purpose="Set out the delivery controls, five-day diagnostic schedule, client election process and detailed operating model for the six separate scope engagements.",
        meta=[
            ("Prepared for", "Squadron Energy"),
            ("Prepared by", "AuditCo Australia"),
            ("Date", "11 August 2026"),
            ("Status", "Working draft for scope and commercial discussion"),
        ],
    )
    add_heading(doc, "1. Purpose And Use")
    add_para(doc, "This plan describes a rapid diagnostic and a controlled menu of six separate scope engagements. The diagnostic is a standalone service that introduces all six scope areas at an introductory level for five working days. It gives Squadron Energy a cross-scope evidence and risk view and identifies the area that needs detailed focus; it does not complete any scope or provide a final rating.")
    add_para(doc, "After the diagnostic, Squadron Energy may elect a separate detailed engagement for the recommended scope area, another scope area or more than one scope area. Each elected scope area has its own written authorisation, fee, timing, evidence request, work plan, report and action tracker.")
    add_para(doc, "The later scope engagements are deliberately bounded and timeboxed. They provide a scope-specific design assessment and action tracker; they are not portfolio-wide audits or sustained operating-effectiveness tests.")
    add_callout(doc, "Core operating rule", "No separate scope engagement starts automatically. The client chooses whether to activate it, which scope to activate, the order of activation and whether to pause or close future work.", fill=PALE_GREEN, accent=LIME)

    add_heading(doc, "2. Delivery Model")
    add_table(doc, ["Stage", "Client Decision", "AuditCo Work", "Output", "Commercial Treatment"], [
        ("Rapid Diagnostic", "Authorise the five-day cross-scope diagnostic.", "Complete an introductory scan across the six scope areas over five working days.", "Diagnostic brief, priority scope recommendation and decision options.", "Fixed fee of $6,000 excluding goods and services tax."),
        ("Client Election", "Select the next scope engagement, wait or close.", "Prepare the selected scope confirmation and finalise assumptions.", "Written election or agreed authorisation.", "Separate fee confirmation."),
        ("Scope Engagement", "Approve one scope area to start.", "Complete a bounded, timeboxed review and reporting.", "Scope-specific final report and action tracker.", "Separate fixed or capped fee."),
        ("Further Engagement", "Select another scope area, pause or stop.", "Repeat the controlled scope process.", "New scope report or programme status update.", "New separate fee confirmation."),
    ], [1650, 2350, 3000, 1800, 1090], font_size=7.5, compact=True)

    add_heading(doc, "3. Scope Menu")
    add_scope_focus_table(doc, SCOPES)
    add_heading(doc, "Indicative Timing By Scope Area", 2)
    add_scope_timing_summary_table(doc, SCOPES)
    add_callout(doc, "Timeboxed scope model", "Each separately elected scope engagement is planned at 3-5 working days: three days for a focused review, four days for the standard basis and five days for greater complexity. The exact duration and fee are confirmed after the client election and evidence scan.", fill=PALE_ORANGE, accent=ORANGE, size=9.3)

    add_heading(doc, "4. Diagnostic Readiness Gate")
    add_para(doc, "AuditCo will complete the following checks before Day 1. If a readiness item is incomplete, the condition will be recorded and the parties will decide whether to start with a narrowed boundary, change the evidence plan or pause.")
    add_table(doc, ["Readiness Item", "Minimum Position", "Evidence Of Readiness"], [
        ("Diagnostic Boundary", "The six scope areas, diagnostic objective, key assets and lifecycle stages, and exclusions are agreed.", "Written diagnostic confirmation."),
        ("Sponsor", "One accountable sponsor and one coordinator are nominated.", "Contact and escalation list."),
        ("Evidence Access", "The evidence channel and document permissions work.", "Access test and evidence index."),
        ("Initial Evidence", "The available evidence index or diagnostic evidence set is available, or its absence is recorded.", "Evidence tracker."),
        ("People", "Priority scope-owner interviews and, where useful, one short risk-pathway walkthrough are available.", "Interview and validation schedule."),
        ("Criteria", "Design criteria and assurance boundary are agreed.", "Criteria note."),
        ("Response Windows", "Client confirms when it can answer questions and correct facts.", "Response plan."),
        ("Independence", "AuditCo confirms independence, conflicts and required capability.", "Independence and competency record."),
    ], [2200, 5000, 2690], font_size=7.9, compact=True)

    add_heading(doc, "5. Five-Day Diagnostic Schedule")
    add_table(doc, ["Day", "Focus", "AuditCo Work", "Squadron Energy Input", "Daily Output"], [
        ("Day 1", "Mobilise And Map", "Hold start meeting. Confirm the six scope areas, diagnostic boundary, key assets and lifecycle stages, criteria, priority risk questions, evidence index, people and communication route. Issue cross-scope evidence request.", "Confirm the diagnostic boundary, sponsor, coordinator, scope owners, available evidence and interview availability.", "Approved diagnostic plan and cross-scope evidence request."),
        ("Day 2", "Review Evidence And Criteria", "Review framework, obligations, risks, controls, roles, reporting and selected records for each scope area. Build a cross-scope readiness matrix.", "Provide controlled documents or an evidence index and explain missing or superseded records.", "Cross-scope evidence and readiness matrix."),
        ("Day 3", "Targeted Interviews And Validation", "Complete priority interviews across scope owners and, where useful, one short walkthrough of the most material risk pathway. Validate the highest-priority questions.", "Provide attendees, access and process information; identify contrary evidence and practical constraints.", "Validated priority questions and updated evidence log."),
        ("Day 4", "Prioritise And Challenge", "Compare scope areas by risk exposure, evidence position, apparent design concern, change and management priority. Apply high-level MATE and CONTRA challenge.", "Review clarification questions and provide missing records or explain limitations.", "Priority scope recommendation and decision options."),
        ("Day 5", "Brief And Election Pack", "Issue diagnostic brief with recommended scope, rationale, evidence limitations, suggested procedures and separate engagement proposal.", "Confirm factual corrections and decide whether to request a separate scope proposal or wait.", "Diagnostic brief and client election pack."),
    ], [850, 1850, 3300, 2500, 1390], font_size=7.35, compact=True)
    add_callout(doc, "Diagnostic output", "The brief will distinguish evidence-supported observations from matters needing further work. It will set out the six scope areas considered, exclusions, evidence received and gaps, interviews and validation completed, cross-scope MATE observations, contrary evidence, limitations, priority questions, the recommended scope and next scope options.", fill=PALE_BLUE, accent=BLUE, size=9.3)

    add_heading(doc, "6. Diagnostic Workpapers")
    add_bullets(doc, [
        "Diagnostic Scope Note And Cross-Scope Boundary.",
        "Evidence Index And Readiness Record.",
        "Cross-Scope Evidence And Risk Triage.",
        "Interview And Walkthrough Evidence Log.",
        "Cross-Scope MATE Screen.",
        "CONTRA Evidence Challenge Record.",
        "Evidence Limitation Register.",
        "Diagnostic Decision Record.",
        "Diagnostic Brief.",
    ], size=9.6, after=2.5)
    add_para(doc, "The workpapers will show what AuditCo reviewed, what it did not review and how each observation was formed.")

    add_heading(doc, "7. Client Election And Activation")
    add_table(doc, ["Election Step", "Required Decision Or Action", "Record"], [
        ("Review the diagnostic brief", "Confirm factual corrections and consider the apparent priority questions.", "Diagnostic factual response."),
        ("Choose a scope", "Select the recommended scope, another area, more than one area or no further work.", "Client election record."),
        ("Confirm scope basis", "Agree objective, exclusions, assets, lifecycle stages, criteria, procedures and evidence assumptions.", "Separate scope confirmation."),
        ("Approve commercial terms", "Approve the separate fixed or capped fee, timing, payment points and assumptions.", "Written authorisation."),
        ("Activate", "Nominate sponsor, coordinator, evidence owners and authorised participants.", "Gate 0 - Start readiness."),
    ], [2200, 5400, 2290], font_size=8.0, compact=True)
    add_para(doc, "Diagnostic evidence can inform a later engagement, but AuditCo will confirm what can be reused and whether it is sufficient for the elected scope. The client may activate one engagement, several engagements or none.")

    add_heading(doc, "8. Standard Delivery Plan For Each Scope Engagement")
    add_standard_scope_sequence(doc)
    add_callout(doc, "Timing position", "Each separate engagement is planned at 3-5 working days. The final duration depends on the authorised boundary, evidence readiness, interviews, walkthroughs, sample depth, technical specialist input and reporting needs. AuditCo confirms the work sequence, client response windows, expected issue point and fee before activation. There is no automatic combined programme timing commitment.", fill=PALE_ORANGE, accent=ORANGE, size=9.4)

    add_heading(doc, "9. Common Scope Engagement Records")
    add_bullets(doc, [
        "Scope Confirmation.",
        "Criteria-To-Control Matrix.",
        "Connected Assurance Register.",
        "Asset And Hazard Coverage Matrix.",
        "Sample And Deviation Register.",
        "Interview And Walkthrough Evidence Logs.",
        "Evidence Limitation Register.",
        "MATE Assessment Record.",
        "CONTRA Evidence Challenge Record.",
        "Factual Review Tracker.",
        "Finding-To-Action Trace.",
        "Final Report And Action Tracker.",
    ], size=9.4, after=2.2)

    add_heading(doc, "10. Common Scope Engagement Report")
    add_bullets(doc, [
        "Executive summary.",
        "Scope, criteria and assurance boundary.",
        "Assets, lifecycle stages, sample and limitations.",
        "Method and evidence confidence.",
        "Control Design Adequacy assessment.",
        "Findings and supporting evidence.",
        "Contrary evidence and unresolved matters.",
        "Systemic causes or clearly labelled management hypotheses.",
        "Practical recommendations.",
        "Management actions, owners, dates and dependencies.",
        "Contractor or project observations where relevant.",
        "Committee decisions requested, where separately agreed.",
    ], size=9.4, after=2.2)
    add_callout(doc, "Assurance boundary", "Operational Effectiveness will be shown as Not Assessed unless the authorised scope expressly includes broader testing. The report will not provide legal advice, engineering certification or contractor-performance assurance.", fill=PALE_BLUE, accent=BLUE, size=9.3)

    add_heading(doc, "11. Programme Scope Tracker")
    add_table(doc, ["Field", "Purpose"], [
        ("Scope Area", "Identifies the available scope."),
        ("Diagnostic Status", "Shows whether the diagnostic has considered the area."),
        ("Client Election", "Records whether the client elected detailed work."),
        ("Authorisation", "Records written approval and date."),
        ("Engagement Status", "Shows planned, active, paused, completed or closed."),
        ("Timing And Fee", "Records the confirmed scope timing and approved separate fee."),
        ("Evidence Position", "Shows available, missing and restricted evidence."),
        ("Report Status", "Shows draft, factual review, final or issued."),
        ("Actions And Next Decision", "Records owners, dates, dependencies and the next client election or close decision."),
    ], [2700, 7190], font_size=8.2, first_col_fill=PALE_NAVY, compact=True)
    add_para(doc, "The tracker supports clear commercial and governance decisions. It does not create an obligation to activate a scope.")

    add_heading(doc, "12. Change Control And Close")
    add_bullets(doc, [
        "Record and approve changes to scope area, objective, assets, entities, jurisdictions or lifecycle stages before affected work proceeds.",
        "Record and approve changes to evidence, interviews, walkthroughs, samples, technical specialist input or deliverables before affected work proceeds.",
        "Record the effect of an approved change on evidence confidence, findings, actions, timing and fee.",
        "At close, issue the final report and action tracker, transfer agreed actions, record unresolved limitations and update the programme scope tracker.",
        "Ask whether Squadron Energy wants another scope proposal. Pause or close when the client decides.",
    ], size=9.5, after=2.5)
    add_para(doc, "The programme has no fixed end point. Each further scope engagement starts only after a new client election and scope confirmation.", size=9.5, italic=True, color=MID_GREY, after=0)

    path = OUT_DIR / "02_AuditCo_Squadron_Energy_Rapid_Diagnostic_and_Scope_Engagement_Delivery_Plan.docx"
    doc.save(path)
    return path


def build_scope_plan(scope: dict) -> Path:
    doc = Document()
    configure_document(doc)
    add_title_block(
        doc,
        title=f"{scope['name']} - Separate Engagement Plan",
        subtitle="Scope-specific delivery plan activated only by client election",
        purpose=scope["objective"],
        meta=[
            ("Prepared for", "Squadron Energy"),
            ("Prepared by", "AuditCo Australia"),
            ("Commercial basis", "Separate fixed or capped fee confirmed before activation"),
            ("Timing basis", "Confirmed after election, evidence scan and scope confirmation"),
        ],
    )
    add_heading(doc, "1. Scope Objective And Boundary")
    add_para(doc, scope["objective"])
    add_para(doc, "This plan is a controlled work package. It does not start automatically after the diagnostic and does not commit Squadron Energy to activate this scope. The authorised scope confirmation will identify the selected assets, entities, jurisdictions, lifecycle stages, exclusions, reporting audience and evidence assumptions.")
    add_callout(doc, "Activation rule", "Squadron Energy may elect this scope after the diagnostic, elect it before another scope, pause it or decide not to proceed. AuditCo will confirm the fee and expected issue point before work starts.", fill=PALE_GREEN, accent=LIME, size=9.7)

    add_heading(doc, "2. Key Questions")
    add_bullets(doc, scope["questions"], size=9.8, after=2.5)

    add_heading(doc, "3. Evidence Request")
    add_bullets(doc, scope["evidence"], size=9.6, after=2.3)
    add_para(doc, "The evidence list is a planning basis. The final request will be narrowed or expanded through the scope confirmation and evidence-readiness gate. Missing, restricted, superseded or contradictory evidence will be recorded and its effect on confidence will be stated.", size=9.3, italic=True, color=MID_GREY)

    add_heading(doc, "4. Delivery Sequence")
    add_standard_scope_sequence(doc)
    add_heading(doc, "Indicative Timing For This Scope", 2)
    add_table(doc, ["Basis", "Timeboxed Delivery"], [
        ("Focused - 3 working days", scope["timing"]["focused"].split(" - ", 1)[1]),
        ("Standard - 4 working days", scope["timing"]["standard"].split(" - ", 1)[1]),
        ("Complex - 5 working days", scope["timing"]["complex"].split(" - ", 1)[1]),
    ], [2500, 7390], font_size=8.2, first_col_fill=PALE_NAVY, compact=True)
    add_callout(doc, "Timing checkpoint", "This is a standalone 3-5 working-day engagement. The final duration and fee are confirmed after the client election, evidence scan and scope planning discussion. Client response gaps or approved scope changes may move the issue point.", fill=PALE_ORANGE, accent=ORANGE, size=9.3)

    add_heading(doc, "5. Scope Procedures")
    add_numbered(doc, scope["procedures"], size=9.5, after=2.5)

    add_heading(doc, "6. Interviews And Stakeholders")
    add_bullets(doc, scope["interviews"], size=9.5, after=2.2)
    add_para(doc, "The final interview list will be risk-based and will include only people relevant to the authorised boundary. Squadron Energy will nominate attendees and confirm availability before the planned fieldwork starts.", size=9.3, italic=True, color=MID_GREY)

    add_heading(doc, "7. Walkthroughs And Samples")
    add_bullets(doc, scope["walkthroughs"], size=9.5, after=2.2)
    add_para(doc, "Walkthroughs and samples are bounded evidence checks. They are used to understand and corroborate control design; they do not provide assurance over every asset, project, event, contractor or operating condition unless separately authorised.", size=9.3, italic=True, color=MID_GREY)

    add_heading(doc, "8. Scope Outputs And Acceptance")
    add_table(doc, ["Output", "Acceptance Evidence"], [
        ("Scope Confirmation", "Written election, objective, boundary, criteria, evidence assumptions, fee and timing recorded."),
        ("Criteria-To-Control Matrix", "Material criteria, risks, controls, owners, triggers, escalation routes and evidence linked."),
        ("Evidence And Limitation Register", "Evidence reviewed, missing or restricted items, contrary evidence and limitations recorded."),
        ("Assessment And Challenge Record", "MATE assessment and CONTRA challenge completed for material control families."),
        ("Scope-Specific Report", "Findings, evidence confidence, limitations, recommendations and Control Design Adequacy conclusion issued."),
        ("Action Tracker", "Agreed actions, owners, dates, priorities, dependencies and governance handover recorded."),
    ], [3000, 6890], font_size=8.1, first_col_fill=PALE_NAVY, compact=True)
    add_para(doc, "Scope-specific outputs:", bold=True, size=9.8, before=2, after=2)
    add_bullets(doc, scope["outputs"], size=9.4, after=2.0)

    add_heading(doc, "9. Timing, Fee And Change Control")
    add_para(doc, "The engagement is expected to be delivered in 3-5 working days. Three days is a focused review, four days is the standard basis and five days is used for greater complexity. The final duration depends on the authorised boundary, evidence volume and readiness, interview and walkthrough requirements, sample depth, technical specialist input, reporting needs and client response windows. AuditCo will confirm the work sequence and expected issue point before activation.")
    add_para(doc, "The fee will be a separate fixed or capped fee agreed for this scope. Travel, specialist input, additional sampling, extra report rounds or other disbursements outside the confirmed basis require prior written approval.")
    add_para(doc, "Changes to scope, assets, lifecycle stages, evidence, interviews, walkthroughs, samples, deliverables, timing, response windows or fee will be documented and approved before the affected work proceeds.")

    add_heading(doc, "10. Assurance Boundary")
    add_bullets(doc, [
        "The review assesses control design and limited evidence of implementation to corroborate the design assessment.",
        "Operational Effectiveness is Not Assessed unless broader testing is expressly authorised in the scope confirmation.",
        "The engagement does not provide legal advice, engineering certification, site compliance certification or portfolio-wide contractor-performance assurance.",
        "Any evidence limitation or approved deviation that could affect the conclusion will be stated in the report.",
    ], size=9.4, after=2.3)
    add_callout(doc, "Close and next election", "At close, AuditCo will issue the final report and action tracker, transfer agreed actions, record unresolved evidence limitations and ask whether Squadron Energy wants another separate scope proposal.", fill=PALE_BLUE, accent=BLUE, size=9.4)

    path = OUT_DIR / f"{scope['file']}.docx"
    doc.save(path)
    return path


def build_infographic() -> Path:
    doc = Document()
    configure_document(doc, infographic=True)
    add_para(doc, "Squadron Energy", size=9.5, bold=True, color=BLACK, after=1, keep=True)
    add_para(doc, "SQE Delivery Journey", size=21, bold=True, color=NAVY, after=1, line_spacing=0.95, keep=True)
    add_para(doc, "One rapid diagnostic. Client-elected, separately billed scope engagements.", size=10.3, bold=True, color=DARK_BLUE, after=5, line_spacing=1.0, keep=True)

    add_callout(doc, "The model", "Start with a five-working-day diagnostic that introduces all six scope areas at a high level. Use the diagnostic brief to identify which scope needs detailed focus. Every later scope is a separate client election, fee, timing and engagement.", fill=PALE_BLUE, accent=BLUE, size=8.7)

    step_rows = [
        ("1", "Map", "Introduce the six scope areas and confirm priority risks, evidence and stakeholders."),
        ("2", "Diagnose", "Five working days: cross-scope evidence scan, targeted discussions, risk prioritisation and challenge."),
        ("3", "Decide", "Receive the diagnostic brief: cross-scope evidence position, apparent gaps, limitations and recommended scope."),
        ("4", "Elect", "Choose the recommended scope, another area, more than one area, or pause. Confirm written authorisation, fee and timing."),
        ("5", "Deliver", "Complete the elected scope engagement and issue the scope-specific report and action tracker."),
    ]
    add_table(doc, ["Step", "Journey", "What Happens"], step_rows,
              [700, 1900, 7290], font_size=8.0, first_col_fill=PALE_GREEN,
              alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT], compact=True)

    add_table(doc, ["Client Can Activate", "Client Can Also"], [
        ("The recommended scope area", "Choose another scope area first"),
        ("A further scope engagement", "Activate more than one scope through separate confirmations"),
        ("A new scope when ready", "Pause or close without activating every available area"),
    ], [4945, 4945], font_size=8.0, header_fill=DARK_BLUE, first_col_fill=PALE_NAVY, compact=True)

    add_heading(doc, "Six Scope Areas", 2)
    scope_rows = [
        ("Governance And Oversight", "Authority, reporting and escalation"),
        ("Safety Management Framework", "Framework, roles, competence and assurance"),
        ("Risk Register Integration", "Risks, controls, owners and review triggers"),
        ("Incident Learning", "Investigation, learning, actions and improvement"),
        ("Third-Party Governance", "Contractors, critical controls and retained duties"),
        ("Safety In Design And Handover", "Design risk, residual risk and operational transfer"),
    ]
    add_table(doc, ["Scope Area", "Focus"], scope_rows,
              [5100, 4790], font_size=7.9, first_col_fill=PALE_GREEN, compact=True)
    add_callout(doc, "Commercial guardrails", "Diagnostic: standalone fixed fee. Later scopes: separate fixed or capped fee and confirmed timing. No scope starts without written election or authorisation. Timing depends on the elected scope and evidence position.", fill=PALE_ORANGE, accent=ORANGE, size=8.7)
    add_para(doc, "Journey outcome: a clear, evidence-based next decision after each engagement.", size=9.0, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)

    path = OUT_DIR / "09_AuditCo_Squadron_Energy_Delivery_Journey_Infographic_A4.docx"
    doc.save(path)
    return path


def build_manager_email() -> Path:
    email = """Subject: SQE pack update - rapid diagnostic and elected scope engagements

Hi Chris,

I have updated the SQE EOI and delivery pack and saved the files in:

SQE > 05 Rapid Diagnostic and Scope Engagement Pack

The main change is the delivery model. It now separates the rapid diagnostic from any later detailed scope work:

- The initial service is a rapid five-working-day diagnostic across all six scope areas at an introductory level.
- The diagnostic gives a cross-scope evidence position, priority risk questions and a recommendation on which scope to focus on next.
- It does not complete any scope area or provide a final rating.
- Squadron Energy activates any later scope only when it chooses to do so.
- Each later scope is a separate billed engagement with its own authorisation, fee, timing, evidence request, work plan, report and action tracker.
- Indicative delivery is 3-5 working days per elected scope: three days focused, four days standard and five days complex.
- The final timing and fee are confirmed before each scope starts.
- The current draft includes a $6,000 diagnostic fee excluding goods and services tax. Later scopes have separate fixed or capped fees.
- There is no automatic commitment to activate all six scope areas.

This position reflects the difficulty of completing a meaningful detailed scope review within the five-day diagnostic. The diagnostic is an introduction and decision point that identifies the best next focus. It is not a compressed full-scope review.

The folder includes the revised EOI, master delivery plan, six standalone scope plans, simplified A4 infographic and Word and PDF versions.

Please review the commercial position, scope boundaries, timing basis, team and technical specialist assumptions, and independence requirements before external issue.

Kind regards,

Alan
"""
    path = OUT_DIR / "Manager_Email_Copy_Paste.md"
    path.write_text(email, encoding="utf-8")
    return path


def main() -> None:
    outputs: list[Path] = []
    outputs.append(build_eoi())
    outputs.append(build_master_plan())
    for scope in SCOPES:
        outputs.append(build_scope_plan(scope))
    outputs.append(build_infographic())
    outputs.append(build_manager_email())
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
