from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.table import Table


ROOT = Path(r"LOCAL_HOME\Documents\agentic-os-workspace\sqe")
EOI_SOURCE = ROOT / "260723 EOI and Plan" / "AuditCo_Squadron_Energy_EOI_Current.docx"
PLAN_SOURCE = ROOT / "260723 EOI and Plan" / "AuditCo_Squadron_Energy_Delivery_Plan_Current.docx"
EOI_OUTPUT = ROOT / "260723 EOI and Plan" / "AuditCo_Squadron_Energy_EOI_Phased_Diagnostic_Selection.docx"
PLAN_OUTPUT = ROOT / "260723 EOI and Plan" / "AuditCo_Squadron_Energy_Delivery_Plan_Phased_Diagnostic_Selection.docx"


def all_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def find_paragraph(document, text):
    for paragraph in all_paragraphs(document):
        if text in paragraph.text:
            return paragraph
    raise ValueError(f"Could not find paragraph containing: {text}")


def replace_paragraph_text(paragraph, text):
    run_properties = None
    if paragraph.runs:
        source_properties = paragraph.runs[0]._element.rPr
        if source_properties is not None:
            run_properties = deepcopy(source_properties)
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    new_run = paragraph.add_run(text)
    if run_properties is not None:
        new_run._element.insert(0, run_properties)


def set_labeled_paragraph(paragraph, label, body):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    label_run = paragraph.add_run(label)
    label_run.bold = True
    paragraph.add_run(body)


def set_cell_text(cell, text):
    if not cell.paragraphs:
        cell.add_paragraph()
    replace_paragraph_text(cell.paragraphs[0], text)
    for paragraph in list(cell.paragraphs[1:]):
        paragraph._element.getparent().remove(paragraph._element)


def set_row_text(row, values):
    if len(row.cells) != len(values):
        raise ValueError("Row width does not match supplied values")
    for cell, value in zip(row.cells, values):
        set_cell_text(cell, value)


def insert_styled_row(table, row_index, values):
    template = table.rows[row_index]._tr
    new_row = deepcopy(template)
    xml_children = list(table._tbl)
    template_position = xml_children.index(template)
    table._tbl.insert(template_position + 1, new_row)
    set_row_text(table.rows[row_index + 1], values)


def append_styled_row(table, template_index, values):
    template = table.rows[template_index]._tr
    new_row = deepcopy(template)
    table._tbl.append(new_row)
    set_row_text(table.rows[-1], values)


def move_xml_after(anchor_xml, new_xml):
    parent = new_xml.getparent()
    if parent is not None:
        parent.remove(new_xml)
    anchor_xml.addnext(new_xml)


def insert_paragraph_after(document, anchor, text, style="Normal"):
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(text)
    move_xml_after(anchor._p if hasattr(anchor, "_p") else anchor, paragraph._p)
    return paragraph


def clone_table_after(document, anchor, source_table, header_values, data_rows):
    cloned_xml = deepcopy(source_table._tbl)
    for row_xml in list(cloned_xml.tr_lst)[1:]:
        cloned_xml.remove(row_xml)

    cloned_table = Table(cloned_xml, document)
    set_row_text(cloned_table.rows[0], header_values)
    template_row = source_table.rows[1]._tr
    for values in data_rows:
        cloned_xml.append(deepcopy(template_row))
        set_row_text(cloned_table.rows[-1], values)

    move_xml_after(anchor._p if hasattr(anchor, "_p") else anchor, cloned_xml)
    return cloned_table


def add_paragraph_sequence(document, anchor, blocks):
    current = anchor
    for text, style in blocks:
        current = insert_paragraph_after(document, current, text, style)
    return current


def revise_eoi():
    copy2(EOI_SOURCE, EOI_OUTPUT)
    document = Document(EOI_OUTPUT)

    replacements = {
        "AuditCo proposes a seven-week internal audit of the design of Squadron Energy's safety risk management framework for a fixed fee of $37,800 excluding goods and services tax. This is the recommended and conforming base proposal.":
            "AuditCo offers two delivery routes. The direct route is a seven-week internal audit of all six scope areas for a fixed fee of $37,800 excluding goods and services tax. The phased route starts with a five-working-day diagnostic across all six scope areas, then completes one elective focus area per authorised engagement. The direct route remains available when Squadron Energy wants all six areas reviewed in one programme.",
        "The base proposal is a seven-week review with defined mobilisation, evidence, assessment, reporting and committee-presentation stages. A six-week route is available only if the stated readiness conditions are met.":
            "The direct route is a seven-week review with defined mobilisation, evidence, assessment, reporting and committee-presentation stages. A six-week route is available only if the stated readiness conditions are met. The phased route uses a five-working-day diagnostic, a post-diagnostic selection gate and one focus area per authorised engagement.",
        "The seven-week review is offered for $37,800 excluding goods and services tax. Optional services are separately identified.":
            "The direct seven-week review is offered for $37,800 excluding goods and services tax. The phased route is separately priced by phase and focus area.",
        "Other Delivery Choices": "Phased Delivery Route And Other Choices",
        "The planned effort is 32-40 consulting days. The fee and timing assume timely access to agreed documents, systems and participants.":
            "The planned effort for the direct route is 32-40 consulting days. The fee and timing assume timely access to agreed documents, systems and participants. The phased route has a separate allowance for each phase and focus area.",
        "Five-Day Diagnostic. For $6,000 excluding goods and services tax, AuditCo can provide a rapid evidence-readiness and design diagnostic using a limited document set, up to six interviews and one walkthrough. It will identify priority gaps and recommended next steps, but it will not provide a final rating or assurance conclusion.":
            "Five-Day Diagnostic. For $6,000 excluding goods and services tax, AuditCo will complete a risk-based scan across all six scope areas. The work will use a limited document set, up to six priority interviews and one risk-selected walkthrough. It will identify apparent gaps, evidence limits and the scope area that would benefit most from detailed review. It will not complete any scope area or provide a final rating or assurance conclusion.",
        "Optional Phase 2 - Targeted Validation And Final Assurance. If Squadron Energy proceeds after the diagnostic, AuditCo would reuse suitable diagnostic work, close the remaining evidence gaps and complete the procedures needed for a final control-design assurance conclusion. The scope and fixed or capped fee are to be confirmed before this phase starts.":
            "Phase 2 - One Elective Focus Area. After the diagnostic, Squadron Energy will select one focus area from the remaining available scope areas. AuditCo will complete the agreed procedures for that one area and issue a final Control Design Adequacy conclusion for that area. The Phase 2 scope, fee or cap, evidence request, sample and timetable will be approved before work starts. The diagnostic does not complete a scope area.",
        "Any change to scope, sample, timing, assumptions, deliverables or fees will be documented and approved before the affected work proceeds.":
            "Subsequent One-Scope Engagements. Squadron Energy may authorise one remaining focus area at a time until all six areas are complete. Each work package will have a new scope confirmation, evidence request, timetable, fixed or capped fee, report and action tracker. Squadron Energy may reorder or pause later work when risk or business needs change. Any change to scope, sample, timing, assumptions, deliverables or fees will be documented and approved before the affected work proceeds.",
        "The fees assume predominantly remote delivery. Reasonable travel or other disbursements will be incurred only with prior approval and charged at cost. Proposal validity, payment milestones and any credit for reusable diagnostic work are to be confirmed. Variations require written agreement before the affected work proceeds.":
            "The fees assume predominantly remote delivery. Reasonable travel or other disbursements will be incurred only with prior approval and charged at cost. Proposal validity, payment milestones and any credit for directly reusable diagnostic work are to be confirmed. Phase 2 and later focus-area engagements require written approval before commencement. Variations require written agreement before the affected work proceeds.",
    }
    for paragraph in list(all_paragraphs(document)):
        for old, new in replacements.items():
            if old in paragraph.text:
                replace_paragraph_text(paragraph, paragraph.text.replace(old, new))

    labelled_paragraphs = {
        "Five-Day Diagnostic.": " For $6,000 excluding goods and services tax, AuditCo will complete a risk-based scan across all six scope areas. The work will use a limited document set, up to six priority interviews and one risk-selected walkthrough. It will identify apparent gaps, evidence limits and the scope area that would benefit most from detailed review. It will not complete any scope area or provide a final rating or assurance conclusion.",
        "Phase 2 - One Elective Focus Area.": " After the diagnostic, Squadron Energy will select one focus area from the remaining available scope areas. AuditCo will complete the agreed procedures for that one area and issue a final Control Design Adequacy conclusion for that area. The Phase 2 scope, fee or cap, evidence request, sample and timetable will be approved before work starts. The diagnostic does not complete a scope area.",
        "Subsequent One-Scope Engagements.": " Squadron Energy may authorise one remaining focus area at a time until all six areas are complete. Each work package will have a new scope confirmation, evidence request, timetable, fixed or capped fee, report and action tracker. Squadron Energy may reorder or pause later work when risk or business needs change. Any change to scope, sample, timing, assumptions, deliverables or fees will be documented and approved before the affected work proceeds.",
    }
    for label, body in labelled_paragraphs.items():
        paragraph = find_paragraph(document, label)
        set_labeled_paragraph(paragraph, label, body)

    deliverables = document.tables[6]
    journey = document.tables[7]
    fees = document.tables[9]

    insert_styled_row(deliverables, 1, ["Diagnostic Brief", "Summarises cross-scope evidence readiness, preliminary observations, priority ranking and the recommended next scope."])
    append_styled_row(deliverables, 4, ["Scope-Specific Final Report", "Provides the final Control Design Adequacy conclusion for one authorised focus area."])
    append_styled_row(deliverables, 5, ["Programme Scope Tracker", "Shows completed, active, deferred and remaining scope areas across the staged programme."])

    insert_styled_row(journey, 3, ["Flex Gate - Choose Next Scope", "Select one elective focus area after the diagnostic", "Diagnostic brief and Phase 2 plan", "Approve one focus area, fee or cap and timetable"])

    set_row_text(fees.rows[3], ["Phase 2 - One Elective Focus Area", "Fixed or capped fee agreed after the diagnostic", "to be confirmed"])
    insert_styled_row(fees, 3, ["Subsequent One-Scope Engagement", "Fixed or capped fee agreed for each remaining area", "to be confirmed"])

    timing_anchor = find_paragraph(document, "The planned effort for the direct route is 32-40 consulting days.")
    insert_paragraph_after(document, timing_anchor, "Phased Route Timetable", "Heading 2")
    timetable_heading = find_paragraph(document, "Phased Route Timetable")
    clone_table_after(
        document,
        timetable_heading,
        document.tables[5],
        ["Phased Step", "Timing", "Main Activity And Output"],
        [
            ["Phase 1 Diagnostic", "Five working days", "Risk-based scan across all six scope areas; up to six interviews; one risk-selected walkthrough; diagnostic brief and preliminary priority view. No final rating."],
            ["Flex Gate", "Two business days after Phase 1", "Squadron Energy selects one elective focus area from the remaining scope areas and approves the Phase 2 scope, fee or cap and evidence plan."],
            ["Phase 2 Focus Area", "10 working days accelerated; 15 working days normal", "Complete final design procedures for one selected scope area and issue a scope-specific report and action tracker."],
            ["Subsequent Focus Areas", "10-15 working days per area", "Repeat for one remaining area at a time until all selected scope areas are complete."],
        ],
    )

    outputs_heading = find_paragraph(document, "Outputs And Governance")
    outputs_heading.paragraph_format.page_break_before = False

    document.save(EOI_OUTPUT)


def revise_delivery_plan():
    copy2(PLAN_SOURCE, PLAN_OUTPUT)
    document = Document(PLAN_OUTPUT)

    options = document.tables[2]
    append_styled_row(options, 1, ["Phased Route - Phase 1 Diagnostic", "Five working days", "Gate 0 confirms the diagnostic evidence and participant access. The diagnostic scans all six scope areas and provides preliminary insight only."])
    append_styled_row(options, 1, ["Phased Route - Phase 2", "10 working days accelerated or 15 working days normal", "Squadron Energy selects one elective focus area after the diagnostic. A new scope, fee or cap and evidence plan are approved."])
    append_styled_row(options, 1, ["Subsequent One-Scope Engagement", "10-15 working days per area", "One remaining scope area is completed per authorised work package until all areas are complete, subject to Squadron Energy priority."])

    boundaries = document.tables[3]
    append_styled_row(boundaries, 1, ["Five-Day Diagnostic", "Scans all six areas and creates a preliminary priority view. It does not provide a final rating or close a scope area."])
    append_styled_row(boundaries, 1, ["One Scope Per Engagement", "Phase 2 and later engagements cover one elective scope area only. Each has its own scope, fee, timetable and report."])
    append_styled_row(boundaries, 1, ["Cumulative Completion", "The staged programme continues through later one-scope engagements until all six areas are complete, unless Squadron Energy defers or closes the route."])

    replace_map = {
        "The seven-week baseline begins only when Gate 0 confirms that scope, critical evidence, access, specialist capability and response windows are sufficiently ready. The six-week option is a conditional target, not the default commitment.":
            "The seven-week direct baseline begins only when Gate 0 confirms that scope, critical evidence, access, specialist capability and response windows are sufficiently ready. The six-week option is a conditional target, not the default commitment. The phased route begins with a five-working-day diagnostic and uses one scope area per authorised engagement.",
        "Additional entities, jurisdictions, samples, site visits, broader implementation testing, formal maturity assessment or additional review rounds require written agreement on scope, timing and fee.":
            "Additional entities, jurisdictions, samples, site visits, broader implementation testing, formal maturity assessment or additional review rounds require written agreement on scope, timing and fee. For the phased route, each additional focus area is a new work package.",
        "Ready: start the seven-week timetable.":
            "Ready: start the authorised route. The direct route starts the seven-week timetable. The phased route starts the five-day diagnostic.",
        "The proposed fixed professional fee for the seven-week baseline is $37,800. Commercial treatment of goods and services tax, disbursements, payment milestones and validity is to be confirmed in the expression of interest or engagement letter.":
            "The proposed fixed professional fee for the seven-week direct baseline is $37,800. The five-day diagnostic is $6,000 excluding goods and services tax. Phase 2 and later one-scope engagements use a fixed or capped fee agreed after scope selection. Commercial treatment of goods and services tax, disbursements, payment milestones and validity is to be confirmed in the expression of interest or engagement letter.",
        "The 32-40 consulting-day allowance assumes:":
            "The 32-40 consulting-day allowance applies to the direct seven-week review and assumes:",
        "The baseline is seven weeks from Gate 0 recording Ready, or Ready With Conditions that do not affect the critical path, to issue of the final report and committee-ready pack.":
            "The direct baseline is seven weeks from Gate 0 recording Ready, or Ready With Conditions that do not affect the critical path, to issue of the final report and committee-ready pack.",
        "The accelerated target is six weeks only when every accelerated condition is met. The contingency is an indicative eight-to-ten-week reforecast where critical information, access, capability or scope conditions are not met.":
            "The direct accelerated target is six weeks only when every accelerated condition is met. The contingency is an indicative eight-to-ten-week reforecast where critical information, access, capability or scope conditions are not met.",
        "Following the scope discussion and initial document inventory, this plan should be finalised and attached to the engagement letter or mobilisation pack as the agreed delivery baseline.":
            "Phased route: five-day diagnostic; select one remaining scope within two business days; Phase 2 takes 10 accelerated or 15 normal working days; each later scope takes 10-15 working days. All six scopes take about 13-19 working weeks of review time, plus decision and response intervals. Finalise and attach this plan after the initial document inventory.",
    }
    for paragraph in list(all_paragraphs(document)):
        for old, new in replace_map.items():
            if old in paragraph.text:
                replace_paragraph_text(paragraph, paragraph.text.replace(old, new))

    gates = document.tables[11]
    anchor = find_paragraph(document, "The final report and committee-ready pack will be issued at Gate 5.")
    blocks = [
        ("Phased Delivery Route", "Heading 2"),
        ("The phased route separates the diagnostic, scope selection and final assurance. The diagnostic scans all six scope areas. It does not complete a scope area or issue a final rating.", "Normal"),
        ("Phase 1 Diagnostic", "Heading 3"),
        ("Phase 1 is a five-working-day diagnostic. It uses a limited evidence set, up to six priority interviews and one risk-selected walkthrough. It identifies evidence limits, apparent gaps and a recommended priority. The walkthrough is selected after the initial scan and does not complete that scope area.", "Normal"),
        ("Flex Gate", "Heading 3"),
        ("Within two business days after the diagnostic briefing, Squadron Energy selects one elective focus area from the remaining available scope areas. AuditCo then issues the scope, evidence request, timetable, team allocation and fixed or capped fee for approval.", "Normal"),
        ("Phase 2 And Later Scope Engagements", "Heading 3"),
        ("Phase 2 covers one focus area only. It produces a final Control Design Adequacy conclusion for that area. Each later engagement covers one remaining area. Each work package starts with a new readiness decision and ends with its own report and action tracker.", "Normal"),
        ("Phased Route Timetable", "Heading 3"),
    ]
    heading = add_paragraph_sequence(document, anchor, blocks)
    clone_table_after(
        document,
        heading,
        document.tables[6],
        ["Stage", "Timing", "What AuditCo Will Do", "Milestone", "Squadron Energy Input"],
        [
            ["Phase 1 Diagnostic", "Five working days", "Scan all six scope areas; review priority evidence; complete up to six interviews and one risk-selected walkthrough; record limitations.", "Diagnostic brief and Flex Gate recommendation", "Provide evidence, interviews and daily clarifications."],
            ["Flex Gate", "Two business days after Phase 1", "Present the evidence-readiness view, preliminary observations and priority ranking.", "Select one elective focus area", "Sponsor and AuditCo Partner approve scope, fee and Phase 2 plan."],
            ["Phase 2 Focus Area", "10 working days accelerated or 15 working days normal", "Complete mapping, interviews, walkthroughs, bounded sample, challenge and reporting for one focus area.", "Gate 5 - final report for selected area", "Provide evidence, factual response and actions."],
            ["Subsequent Focus Area", "10-15 working days per area", "Repeat the scope-specific route for one remaining area. Reuse approved evidence where relevant.", "New Gate 0 and Gate 5 per area", "Select the next area and approve scope and fee."],
        ],
    )

    insert_styled_row(gates, 3, ["Flex", "Post-Diagnostic Scope Decision", "Diagnostic brief, evidence-readiness view, priority ranking, selected one scope area, approved Phase 2 scope and fee.", "Internal Audit Sponsor And AuditCo Engagement Partner"])
    replace_gate = find_paragraph(document, "The direct baseline is seven weeks from Gate 0 recording Ready")
    insert_paragraph_after(document, replace_gate, "The phased route does not replace Gates 0-6. It adds the Flex Gate after the diagnostic. Each later focus-area engagement begins with a new Gate 0 and ends with its own Gate 5 report.", "Normal")

    document.save(PLAN_OUTPUT)


if __name__ == "__main__":
    revise_eoi()
    revise_delivery_plan()
    print(EOI_OUTPUT)
    print(PLAN_OUTPUT)
