from __future__ import annotations

from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree


SOURCE = Path(r"LOCAL_HOME\OneDrive - AuditCo\SQE\SQE-EOI-Review-260810.docx")
OUTPUT = Path(r"LOCAL_HOME\OneDrive - AuditCo\SQE\SQE-EOI-Review-260810-REISSUED.docx")

# Colours sampled from the supplied AuditCo document and its retained theme.
BRAND_BLUE = "275070"
DEEP_NAVY = "0E2841"
BRAND_ORANGE = "FF8200"
PALE_BLUE = "EAF5EF"
PALE_ORANGE = "FFF0E0"
PALE_GREY = "F4F6F8"
MID_GREY = "D9E1E8"
BODY = "202124"
WHITE = "FFFFFF"
BODY_SIZE = 10.5
TABLE_SIZE = 10.0
MINOR_LABEL_WORDS = {"a", "an", "and", "as", "at", "by", "for", "from", "in", "into", "nor", "of", "on", "or", "the", "to", "with"}


def title_case_label(text):
    """Apply readable title case to headings and table labels."""
    words = str(text).split(" ")
    converted = []
    for index, word in enumerate(words):
        if not word:
            converted.append(word)
            continue
        parts = word.split("-")
        converted_parts = []
        for part_index, part in enumerate(parts):
            if not part:
                converted_parts.append(part)
                continue
            if part.isupper() and any(character.isalpha() for character in part):
                converted_parts.append(part)
                continue
            if any(character.isupper() for character in part[1:]):
                converted_parts.append(part)
                continue
            lower = part.lower()
            if index > 0 and part_index == 0 and lower in MINOR_LABEL_WORDS:
                converted_parts.append(lower)
            else:
                converted_parts.append(lower[:1].upper() + lower[1:])
        converted.append("-".join(converted_parts))
    return " ".join(converted)


def remove_children(element, keep_tags=()):
    for child in list(element):
        if child.tag not in keep_tags:
            element.remove(child)


def set_run_font(run, size=BODY_SIZE, bold=False, color=BODY, font_name="Aptos"):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), font_name)


def clear_paragraph(paragraph):
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def set_paragraph(paragraph, text, size=BODY_SIZE, bold=False, color=BODY, align=None, after=3, before=0):
    clear_paragraph(paragraph)
    if align is not None:
        paragraph.alignment = align
    pf = paragraph.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing = 1.05
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return paragraph


def add_body(doc, text, size=BODY_SIZE, after=4):
    p = doc.add_paragraph(style="Normal")
    set_paragraph(p, text, size=BODY_SIZE, after=after)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    size = 16 if level == 1 else 12.5
    set_paragraph(p, title_case_label(text), size=size, bold=True, color=DEEP_NAVY, after=4, before=3)
    p.paragraph_format.keep_with_next = True
    return p


def add_page_break(doc):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


def set_cell_margins(cell, top=60, start=80, bottom=60, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in kwargs[edge]:
                element.set(qn(f"w:{key}"), str(kwargs[edge][key]))


def set_cell_text(cell, text, size=TABLE_SIZE, bold=False, color=BODY, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    p = cell.paragraphs[0]
    clear_paragraph(p)
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    parts = str(text).split("\n")
    for index, part in enumerate(parts):
        if index:
            p.add_run().add_break()
        run = p.add_run(part)
        set_run_font(run, size=size, bold=bold, color=color)


def set_table_width(table, widths):
    table.autofit = False
    total_twips = int(sum(widths) * 1440)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, colour=MID_GREY, size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), colour)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_table(doc, headers, rows, widths, body_size=TABLE_SIZE, first_col_fill=PALE_BLUE, header_fill=BRAND_BLUE):
    body_size = TABLE_SIZE
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, widths)
    set_table_borders(table)
    header = table.rows[0]
    repeat_header(header)
    prevent_row_split(header)
    for cell, text in zip(header.cells, headers):
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell, top=65, bottom=65, start=80, end=80)
        set_cell_text(cell, title_case_label(text), size=body_size, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for col_index, (cell, text) in enumerate(zip(row.cells, values)):
            set_cell_margins(cell, top=60, bottom=60, start=80, end=80)
            if col_index == 0:
                fill = first_col_fill
                bold = True
            else:
                fill = "FFFFFF" if row_index % 2 == 0 else PALE_GREY
                bold = False
            set_cell_shading(cell, fill)
            label = title_case_label(text) if col_index == 0 else text
            set_cell_text(cell, label, size=body_size, bold=bold, color=DEEP_NAVY if col_index == 0 else BODY)
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(0)
    return table


def set_paragraph_callout(paragraph, fill=PALE_ORANGE, border=BRAND_ORANGE):
    p_pr = paragraph._p.get_or_add_pPr()
    for tag in ("w:pBdr", "w:shd"):
        existing = p_pr.find(qn(tag))
        if existing is not None:
            p_pr.remove(existing)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), border)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_note(doc, label, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Pt(8)
    p.paragraph_format.right_indent = Pt(4)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.05
    set_paragraph_callout(p)
    run = p.add_run(title_case_label(label) + "  ")
    set_run_font(run, size=BODY_SIZE, bold=True, color=BRAND_ORANGE)
    run = p.add_run(text)
    set_run_font(run, size=BODY_SIZE, color=BODY)
    return p


def configure_content_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(0.59)
    section.bottom_margin = Inches(0.20)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = RGBColor.from_string(BODY)


def build_content():
    doc = Document()
    configure_content_document(doc)

    # Body Page 1
    p = doc.add_paragraph(style="Heading 1")
    set_paragraph(p, "Rapid Five-Day Diagnostic and Elected Scope Engagements", size=18, bold=True, color=DEEP_NAVY, after=2)
    p.paragraph_format.keep_with_next = True
    p = doc.add_paragraph(style="Normal")
    set_paragraph(p, "Safety Risk Management Internal Audit — Expression of Interest Response", size=11.5, bold=True, color=BRAND_BLUE, after=6)
    add_body(doc, "This Expression of Interest sets out a staged service for Squadron Energy, starting with a rapid diagnostic before any separate client-elected scope engagement.", after=4)

    add_heading(doc, "1. Recommended Delivery Model", 1)
    add_body(doc, "AuditCo proposes a standalone rapid diagnostic lasting five working days. It introduces all six scope areas at a high level, tests the evidence position and priority risk questions, and identifies the scope area that should receive detailed focus.")
    add_body(doc, "The diagnostic does not complete a scope area or issue a final scope conclusion. Squadron Energy can elect the recommended scope, another scope, more than one scope through separate confirmations, or no further work.")
    add_note(doc, "Commercial Control", "No later scope engagement starts automatically. AuditCo confirms the scope, work, evidence, fee and timing after Squadron Energy makes an election.")

    add_table(
        doc,
        ["Stage", "Purpose", "Client Decision"],
        [
            ["Rapid Diagnostic", "Five working days. Introduce all six scope areas, scan the evidence position and test priority risk questions.", "Confirm the diagnostic boundary, people and evidence access."],
            ["Diagnostic Brief", "Summarise readiness, priority concerns, evidence limits and the recommended next scope.", "Elect any scope, request a change, or pause."],
            ["Separate Scope Engagement", "Review one elected scope. The expected duration is 3–5 working days, based on complexity.", "Approve a written scope, fee, timing and evidence request."],
        ],
        [1.35, 3.15, 1.70],
    )
    add_heading(doc, "What The Model Gives Squadron Energy", 2)
    add_table(
        doc,
        ["Benefit", "How It Works"],
        [
            ["Clear Starting Point", "The diagnostic frames the work before detailed effort starts."],
            ["Client Choice", "Squadron Energy controls which scope starts next."],
            ["Traceable Work", "Each engagement has its own scope, evidence, report, fee and decision."],
        ],
        [1.55, 4.65],
    )
    add_page_break(doc)

    # Body Page 2
    add_heading(doc, "2. The Six Scope Areas", 1)
    add_body(doc, "During the diagnostic, AuditCo considers each scope area at a high level and uses the available evidence and targeted discussions to identify the best next focus. The detailed scope, assets, lifecycle stages and sample are confirmed after election.")
    add_table(
        doc,
        ["Scope Area", "Primary Question", "Initial Focus"],
        [
            ["Governance And Oversight", "Does governance provide clear authority, information, oversight and escalation for critical safety risk?", "Authority, reporting, oversight and escalation."],
            ["Safety Management Framework", "Does the framework connect requirements, roles, controls, competence and assurance across the selected lifecycle?", "Framework coherence, competence, change and assurance."],
            ["Risk Register Integration", "Do risk registers connect hazards, controls, owners, triggers, actions and assurance across corporate, project and operational levels?", "Risk ownership, controls, triggers and assurance."],
            ["Incident Learning", "Do incidents and high-potential events lead to sound investigation, learning, action and risk improvement?", "Investigation quality, learning and action verification."],
            ["Third-Party Governance", "Does the contractor lifecycle protect Squadron Energy's retained duties and provide suitable oversight of critical safety risks?", "Contractor lifecycle, critical controls and retained duties."],
            ["Safety In Design And Handover", "Does the design and handover process identify, control and transfer residual safety risk into operations?", "Design risk, residual risk and operational transfer."],
        ],
        [1.60, 2.75, 1.85],
    )
    add_heading(doc, "MATE And CONTRA Assessment Lens", 2)
    add_table(
        doc,
        ["Method Component", "Diagnostic Question"],
        [
            ["MATE — Mandate", "Is the control required, approved and linked to the relevant obligation and risk?"],
            ["MATE — Accountability", "Is one named role accountable, with authority and oversight?"],
            ["MATE — Trigger", "What event, threshold or gateway activates the control or review?"],
            ["MATE — Escalation", "Where does unresolved risk go for decision and oversight?"],
            ["CONTRA — Independent Challenge", "What contrary evidence, gaps, limits or alternative explanations could change the view?"],
        ],
        [1.45, 4.75],
    )
    add_page_break(doc)

    # Body Page 3
    add_heading(doc, "3. Scope Area Timing And Delivery", 1)
    add_body(doc, "Each scope area is a separate billed engagement. The timing below is an estimate for the agreed boundary, and AuditCo confirms the duration after Squadron Energy elects the scope and the evidence position is known.")
    add_table(
        doc,
        ["Scope Area", "Summary Breakdown", "Estimated Time"],
        [
            ["Governance And Oversight", "Authority, reporting, oversight and escalation.", "3 days: one governance structure and a limited reporting sample.\n4 days: standard governance and committee review.\n5 days: multiple committees, projects, contractors or complex reporting flows."],
            ["Safety Management Framework", "Framework coherence, competence, change and assurance.", "3 days: one defined framework component.\n4 days: standard framework review.\n5 days: review including Project Kraken, Golden Rules, competence and transition evidence."],
            ["Risk Register Integration", "Risk ownership, controls, triggers and assurance.", "3 days: one risk system or a limited risk sample.\n4 days: standard multi-level risk review.\n5 days: multiple registers, assets, projects or risk taxonomies."],
            ["Incident Learning", "Investigation quality, learning and action verification.", "3 days: one event pathway or a small sample.\n4 days: standard event and action review.\n5 days: multiple events, complex investigations or recurring-theme analysis."],
            ["Third-Party Governance", "Contractor lifecycle, critical controls and retained duties.", "3 days: one contractor and a narrow lifecycle review.\n4 days: standard contractor lifecycle and contract sample.\n5 days: multiple contractors, projects, jurisdictions or technical interfaces."],
            ["Safety In Design And Handover", "Design risk, residual risk and operational transfer.", "3 days: one design-to-handover risk thread.\n4 days: standard project or asset review.\n5 days: multiple projects, assets, design changes or technical interfaces."],
        ],
        [1.48, 1.72, 3.00],
    )
    add_note(doc, "Timing Rule", "Three working days is a focused review. Four working days is the standard basis. Five working days is used for complex evidence, multiple interfaces, broader sampling or specialist input.")

    add_heading(doc, "4. Assurance Method And Boundary", 1)
    add_body(doc, "AuditCo will use MATE and CONTRA in the diagnostic and in each elected scope engagement. MATE tests control design through Mandate, Accountability, Trigger and Escalation, while CONTRA provides independent challenge of the evidence and proposed conclusion.")
    add_table(
        doc,
        ["Method", "How AuditCo Uses It"],
        [
            ["MATE", "Tests whether each material control family is required, owned, activated and escalated through Mandate, Accountability, Trigger and Escalation."],
            ["CONTRA", "Challenges supporting and contrary evidence, gaps, limitations, alternative explanations and the evidence-to-finding pathway."],
        ],
        [1.55, 4.65],
    )
    add_body(doc, "In the five-day diagnostic, AuditCo applies MATE and CONTRA at a high level across all six scope areas. For an elected scope engagement, the team applies them in detail before drafting findings and recommendations.", after=2)
    add_body(doc, "The review assesses control design and limited evidence of implementation needed to support that assessment. It does not provide legal advice, engineering certification, site compliance certification or sustained operating-effectiveness testing unless a separate scope expressly authorises it.", after=2)
    add_note(doc, "Method Boundary", "MATE and CONTRA support the auditor's assessment. They do not replace the agreed criteria, professional judgement or client review.")
    add_note(doc, "Renewable-Energy Risk Lens", "Where relevant, the criteria may consider high-voltage energisation, arc flash, isolation, stored energy, simultaneous operations, battery hazards, working at height, fatigue and lone work.")
    add_page_break(doc)

    # Body Page 4
    add_heading(doc, "5. Rapid Five-Day Diagnostic", 1)
    add_body(doc, "The diagnostic introduces all six scope areas. It is not a detailed review of one selected scope; instead, it gives Squadron Energy enough information to choose where to focus next.")
    add_table(
        doc,
        ["Day", "Focus", "AuditCo Work", "Squadron Energy Input", "Daily Output"],
        [
            ["Day 1", "Mobilise And Map", "Confirm the six scope areas, boundary, key assets, lifecycle stages, criteria, priority questions, evidence index, people and communication route. Issue the evidence request.", "Confirm the sponsor, coordinator, scope owners, available evidence and interview availability.", "Approved diagnostic plan and cross-scope evidence request."],
            ["Day 2", "Review Evidence And Criteria", "Review available framework, obligations, risks, controls, roles, reporting and selected records for each scope. Build the readiness matrix.", "Provide controlled documents or an evidence index. Explain missing or superseded records.", "Cross-scope evidence and readiness matrix."],
            ["Day 3", "Targeted Interviews And Validation", "Complete priority interviews and, where useful, one short walkthrough of the most material risk pathway. Validate priority questions.", "Provide attendees, access and process information. Identify contrary evidence and constraints.", "Validated priority questions and updated evidence log."],
            ["Day 4", "Prioritise And Challenge", "Compare scope areas by risk exposure, evidence position, apparent design concern, change and management priority. Apply high-level MATE and CONTRA challenge.", "Review clarification questions. Provide missing records or explain limitations.", "Priority scope recommendation and decision options."],
            ["Day 5", "Brief And Election Pack", "Issue the diagnostic brief with the recommended scope, rationale, evidence limitations, suggested procedures and separate engagement proposal.", "Confirm factual corrections. Decide whether to request a separate scope proposal or wait.", "Diagnostic brief and client election pack."],
        ],
        [0.55, 1.25, 2.05, 1.45, 0.90],
    )
    add_note(doc, "Diagnostic Boundary", "The diagnostic does not complete any scope area, provide a final Control Design Adequacy rating, test sustained operating effectiveness, or certify legal or technical compliance.")
    add_heading(doc, "Diagnostic Outputs", 2)
    add_table(
        doc,
        ["Output", "Purpose"],
        [
            ["Evidence-Readiness View", "Identifies available, missing, inconsistent or inaccessible evidence and its effect on the planned work."],
            ["Priority Scope Recommendation", "Shows which scope appears to need detailed focus and why."],
            ["Diagnostic Brief", "Summarises priority questions, apparent gaps, evidence limits and decision options."],
            ["Election Pack", "Sets out the proposed scope, timing, fee basis, evidence request and next engagement steps."],
        ],
        [1.75, 4.45],
    )
    add_page_break(doc)

    # Body Page 5
    add_heading(doc, "6. Separate Scope Engagement Activation", 1)
    add_body(doc, "After the diagnostic, Squadron Energy decides whether to activate a separate scope engagement. The recommended scope is an option, not an automatic next phase.")
    add_table(
        doc,
        ["Client Decision", "AuditCo Action", "Output"],
        [
            ["Elect the recommended scope", "Reuse directly relevant evidence where confirmed. Issue the detailed scope confirmation.", "Separate fee, timing, evidence request and scope work plan."],
            ["Elect another scope", "Review the chosen scope assumptions and confirm the relevant procedures and evidence.", "New scope-specific engagement basis."],
            ["Elect more than one scope", "Prepare separate confirmations and sequence the work by client priority and availability.", "Separate billed engagements with clear boundaries."],
            ["Pause or stop", "Record the decision and retain the diagnostic brief for future use.", "No further work starts without a new election."],
            ["Change the boundary", "Reconfirm assets, lifecycle stages, sample, specialists, fee and timing before work starts.", "Approved variation or revised scope confirmation."],
        ],
        [1.55, 2.90, 1.75],
    )
    add_heading(doc, "Activation Controls", 2)
    add_table(
        doc,
        ["Control", "Requirement"],
        [
            ["Written Election", "Squadron Energy confirms the elected scope and authorised contact."],
            ["Scope Confirmation", "AuditCo records the boundary, criteria, sample, exclusions and assumptions."],
            ["Commercial Confirmation", "AuditCo confirms the separate fee basis, timing and approval route."],
            ["Evidence Request", "AuditCo lists required evidence and the date required for the agreed plan."],
            ["Report And Actions", "AuditCo issues the scope report and records agreed actions, owners and due dates."],
        ],
        [1.65, 4.55],
    )
    add_heading(doc, "7. Individual Scope Delivery Plans — Part 1", 1)
    add_body(doc, "Each plan is refined after election. The review uses a focused sample, applies MATE to each material control family, uses CONTRA before findings are drafted, and traces the selected risk pathway from requirement to action.")
    add_table(
        doc,
        ["Scope Area", "Core Review Activities", "Main Inputs", "Timing"],
        [
            ["Governance And Oversight", "Map mandate, delegations, committee roles, reporting, indicators and escalation. Test a small sample of reports and minutes against material safety risks.", "Governance structure, terms of reference, delegations, reports, minutes and escalation records.", "3–5 days"],
            ["Safety Management Framework", "Map obligations, policies, Golden Rules, roles, competence, change and assurance across the selected lifecycle.", "Framework documents, Project Kraken records, Golden Rules, competence records, change records and assurance plans.", "3–5 days"],
            ["Risk Register Integration", "Trace selected risks to controls, owners, triggers, actions and assurance across corporate, project and operational registers. Review consistency.", "Risk registers, critical-control records, risk decisions, action logs, system extracts and review records.", "3–5 days"],
        ],
        [1.25, 2.60, 1.60, 0.75],
    )
    add_page_break(doc)

    # Body Page 6
    add_heading(doc, "7. Individual Scope Delivery Plans — Part 2", 1)
    add_table(
        doc,
        ["Scope Area", "Core Review Activities", "Main Inputs", "Timing"],
        [
            ["Incident Learning", "Review event classification, investigation, root cause, learning, action ownership and closure verification. Compare recurring themes.", "Incident records, investigation reports, learning notices, action registers and close-out evidence.", "3–5 days"],
            ["Third-Party Governance", "Review contractor lifecycle, retained duties, critical-control requirements, interfaces, oversight and handover.", "Prequalification, contracts, contractor plans, onboarding records, SWMS, interface records and assurance results.", "3–5 days"],
            ["Safety In Design And Handover", "Trace design risk, residual risk, change decisions and transfer into operations. Review readiness and acceptance evidence.", "Design-risk registers, design reviews, change records, commissioning records, handover packs and acceptance records.", "3–5 days"],
        ],
        [1.25, 2.60, 1.60, 0.75],
    )
    add_note(doc, "Scope Timing", "The 3–5 working-day estimate depends on scope complexity, evidence readiness, interview and walkthrough needs, sample depth, technical input and reporting needs.")

    add_heading(doc, "8. Outputs And Governance", 1)
    add_table(
        doc,
        ["Deliverable", "Purpose"],
        [
            ["Scope Confirmation", "Defines the elected scope, criteria, boundaries, sample, exclusions, timing, fee and evidence request."],
            ["Evidence And Criteria Register", "Shows the evidence received, evidence limits, criteria used and traceability to the review questions."],
            ["Scope Report", "Summarises the work performed, design assessment, evidence limits, findings and priority actions."],
            ["Action Register", "Records agreed actions, owners, due dates, status and closure evidence."],
            ["Close-Out Record", "Confirms factual validation, final decisions and any follow-on scope option."],
        ],
        [1.75, 4.45],
    )
    add_body(doc, "AuditCo will use a short mobilisation, evidence review, targeted interviews or walkthroughs, analysis, factual validation and final reporting for each elected scope. AuditCo will raise evidence or scope issues early.")
    add_body(doc, "AuditCo will nominate an Engagement Partner, Lead Internal Auditor, Safety Governance Specialist, Audit Manager or Analyst, and technical specialists where needed. The team will be confirmed for the authorised assets and lifecycle stages.")
    add_heading(doc, "Next Step", 2)
    add_body(doc, "Confirm the diagnostic sponsor and coordinator, agree the five-working-day diagnostic dates, provide the available evidence index and confirm the key stakeholders for the six scope areas.")

    # Body Page 7
    add_page_break(doc)
    add_heading(doc, "9. Team", 1)
    add_body(doc, "AuditCo will confirm the team for the authorised scope, assets and lifecycle stages. The team will combine internal audit, safety governance and technical capability where required.")
    p = doc.add_paragraph(style="Normal")
    set_paragraph(p, title_case_label("Table 9 — Proposed Team"), size=11.5, bold=True, color=BRAND_BLUE, after=2)
    add_table(
        doc,
        ["Role", "Required Capability", "Responsibility", "Availability And Effort"],
        [
            ["Engagement Partner", "Senior internal audit and safety-governance leadership.", "Independence, quality, significant judgements and final conclusion.", "To be confirmed."],
            ["Lead Internal Auditor — Alan Richardson", "Work health and safety, project management, construction, training and management-system audit capability.", "Planning, criteria mapping, interviews, evidence assessment, findings and reporting.", "Availability and approximate effort to be confirmed."],
            ["Safety Governance Specialist", "Executive governance, obligations, risk and assurance experience.", "Challenge governance information, accountability, indicators and escalation design.", "To be confirmed."],
            ["Audit Manager Or Analyst", "Internal audit workpaper, analysis and reporting capability.", "Evidence control, sampling records, analysis, quality checks and action tracking.", "To be confirmed."],
            ["Technical Specialist", "High-voltage electrical safety and renewable-project commissioning experience where relevant.", "Challenge asset-specific criteria for energisation, isolation, stored energy, simultaneous operations and handover.", "To be confirmed."],
        ],
        [1.25, 1.75, 1.85, 1.35],
    )
    add_note(doc, "Team Confirmation", "Names, qualifications, availability and approximate effort are confirmed before each separate scope engagement starts.")

    # Body Page 8
    add_page_break(doc)
    add_heading(doc, "10. Fees", 1)
    add_body(doc, "The fee model separates the rapid diagnostic from later scope work. The diagnostic has a fixed fee, while each elected scope engagement has its own fee confirmed before activation.")
    p = doc.add_paragraph(style="Normal")
    set_paragraph(p, title_case_label("Table 10 — Fee Basis"), size=11.5, bold=True, color=BRAND_BLUE, after=2)
    add_table(
        doc,
        ["Service", "Fee Basis", "Fee"],
        [
            ["Rapid Five-Day Diagnostic", "Fixed fee for the cross-scope diagnostic and diagnostic brief.", "$6,000 excluding goods and services tax."],
            ["Elected Scope Engagement", "Separate fixed or capped fee. Confirmed after the diagnostic, scope election and evidence review.", "To be confirmed."],
            ["Additional Services Outside Agreed Scope", "Approved fixed fee or agreed hourly rate. Written approval is required before work starts.", "To be confirmed."],
            ["Travel And Disbursements", "At cost only where required and approved in writing before commitment.", "If required."],
        ],
        [2.10, 2.85, 1.25],
    )
    add_note(doc, "Fee Control", "Later scope work begins only after Squadron Energy elects it and approves the separate fee and timing.")
    add_heading(doc, "Fee Assumptions", 2)
    add_body(doc, "Fees exclude goods and services tax. The 3–5 working-day scope estimate depends on complexity, evidence readiness, interviews, walkthroughs, sample depth, technical input and reporting needs.")
    add_body(doc, "Any change to the scope, sample, timing, assumptions, deliverables or fee is documented and approved before the affected work proceeds.")

    # Return only ordinary body elements. The target document supplies the
    # cover, headers, footers, section breaks and contact page.
    return [deepcopy(child) for child in doc.element.body if child.tag != qn("w:sectPr")]


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS, "dc": "http://purl.org/dc/elements/1.1/", "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"}


def replace_middle_body(document_xml, content_elements):
    root = etree.fromstring(document_xml)
    body = root.find("w:body", namespaces=NS)
    children = list(body)
    start = None
    end = None
    for index, child in enumerate(children):
        if child.tag != qn("w:p"):
            continue
        text = "".join(child.xpath(".//w:t/text()", namespaces=NS))
        if text.strip().replace(" ", "") in {"EOI:SQE", "EOI:SQE"}:
            start = index
            break
    if start is None:
        raise RuntimeError("Could not locate the supplied EOI body start")
    for index in range(start + 1, len(children)):
        child = children[index]
        if child.tag == qn("w:p") and child.xpath("./w:pPr/w:sectPr", namespaces=NS):
            end = index
            break
    if end is None:
        raise RuntimeError("Could not locate the supplied EOI body end")
    for child in children[start:end]:
        body.remove(child)
    section_break = children[end]
    for element in content_elements:
        section_break.addprevious(deepcopy(element))
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def update_core_properties(core_xml):
    root = etree.fromstring(core_xml)
    title = root.find("dc:title", namespaces=NS)
    subject = root.find("dc:subject", namespaces=NS)
    description = root.find("dc:description", namespaces=NS)
    keywords = root.find("cp:keywords", namespaces=NS)
    if title is not None:
        title.text = "Squadron Energy EOI - Rapid Five-Day Diagnostic And Elected Scope Engagements"
    if subject is not None:
        subject.text = "Safety Risk Management Internal Audit"
    if description is not None:
        description.text = "Expression of Interest for a rapid five-working-day diagnostic and separate client-elected scope engagements."
    if keywords is not None:
        keywords.text = "SQE, safety risk management, diagnostic, scope engagement, AuditCo"
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    content = build_content()
    with ZipFile(SOURCE, "r") as source_zip:
        parts = {info.filename: source_zip.read(info.filename) for info in source_zip.infolist()}
        parts["word/document.xml"] = replace_middle_body(parts["word/document.xml"], content)
        if "docProps/core.xml" in parts:
            parts["docProps/core.xml"] = update_core_properties(parts["docProps/core.xml"])
    with ZipFile(OUTPUT, "w", ZIP_DEFLATED) as output_zip:
        for name, data in parts.items():
            output_zip.writestr(name, data)
    print(OUTPUT)
    print("CONTENT_ELEMENTS", len(content))


if __name__ == "__main__":
    main()
